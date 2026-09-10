"""Non-vacuity and regression tests for generic journal layout rules."""

from __future__ import annotations

import hashlib
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import journal_layout as LAYOUT


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class JournalLayoutTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "fixture.docx"
        document = Document()

        table = document.add_table(rows=4, cols=3)
        values = [
            ("变量", "2018—2020", "N"),
            ("收入", "4.52（0.81；N=13,089）", ""),
            ("Overall effect of X on Y through all pathways", "EYxm-Yxm*-Yx*m+Yx*m**Mx-Mx*", "1.20"),
            ("合并区间", "", "0.30"),
        ]
        for row, row_values in zip(table.rows, values):
            for cell, value in zip(row.cells, row_values):
                cell.text = value
        # Exercise gridSpan-aware inheritance: the merged header covers grid 0:2.
        table.cell(0, 0).merge(table.cell(0, 1)).text = "变量"

        diagram = document.add_table(rows=2, cols=2)
        for row in diagram.rows:
            for cell in row.cells:
                cell.text = "婚姻满意度 ↗→↘"
        diagram.rows[1].cells[0].paragraphs[0].paragraph_format.first_line_indent = 420
        document.save(self.source)

        prepared = Document(self.source)
        self.diagram_hash_before = LAYOUT._xml_sha256(prepared.tables[1]._tbl)
        self.apply_report = LAYOUT.apply_table_layout(prepared)
        self.formatted = self.root / "formatted.docx"
        prepared.save(self.formatted)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_alignment_spacing_indent_and_diagram_contract(self) -> None:
        report = LAYOUT.audit_table_layout(self.formatted)
        self.assertTrue(report["all_pass"], report)
        self.assertEqual(report["coverage"]["target_tables"], [0])
        self.assertEqual(report["coverage"]["target_cell_count"], 11)
        self.assertEqual(report["coverage"]["alignment_checked_count"], 11)
        self.assertEqual(self.apply_report["target_cells"], 11)
        self.assertEqual(self.apply_report["checked_cells"], 11)

        document = Document(self.formatted)
        # English prose stays left; formula and numeric strings stay right.
        self.assertEqual(document.tables[0].cell(2, 0).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(document.tables[0].cell(2, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(document.tables[0].cell(1, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(LAYOUT._xml_sha256(document.tables[1]._tbl), self.diagram_hash_before)
        self.assertGreater(report["diagram_tables"][0]["nonzero_first_line_preserved"], 0)

    def _mutate(self, label: str, action) -> dict:
        path = self.root / f"{label}.docx"
        shutil.copy2(self.formatted, path)
        document = Document(path)
        action(document)
        document.save(path)
        return LAYOUT.audit_table_layout(path)

    def test_wrong_alignment_fails(self) -> None:
        report = self._mutate(
            "alignment",
            lambda document: setattr(document.tables[0].cell(1, 1).paragraphs[0], "alignment", WD_ALIGN_PARAGRAPH.CENTER),
        )
        self.assertFalse(report["all_pass"])
        self.assertIn("alignment", report["failed"])

    def test_nonzero_indent_fails(self) -> None:
        def action(document):
            paragraph = document.tables[0].cell(1, 1).paragraphs[0]
            ind = paragraph._p.get_or_add_pPr().find(qn("w:ind"))
            ind.set(qn("w:firstLine"), "120")
        report = self._mutate("indent", action)
        self.assertFalse(report["all_pass"])
        self.assertIn("indent", report["failed"])

    def test_wrong_spacing_fails(self) -> None:
        def action(document):
            paragraph = document.tables[0].cell(1, 1).paragraphs[0]
            spacing = paragraph._p.get_or_add_pPr().find(qn("w:spacing"))
            spacing.set(qn("w:line"), "999")
        report = self._mutate("spacing", action)
        self.assertFalse(report["all_pass"])
        self.assertIn("spacing", report["failed"])

    def test_mutation_restore_is_byte_identical(self) -> None:
        baseline = sha256(self.formatted)
        path = self.root / "restore.docx"
        shutil.copy2(self.formatted, path)
        document = Document(path)
        document.tables[0].cell(1, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.save(path)
        self.assertNotEqual(sha256(path), baseline)
        shutil.copy2(self.formatted, path)
        self.assertEqual(sha256(path), baseline)

    def test_object_only_page_gate_fails(self) -> None:
        report = LAYOUT.evaluate_object_pages({3, 5}, {3: 80, 5: 0})
        self.assertFalse(report["all_pass"])
        self.assertEqual(report["object_only_pages"], [5])

    def test_object_scope_resolves_exact_page_interval(self) -> None:
        pages = {1: "前言", 2: "（二）基础回归", 3: "表格", 4: "（三）后续分析"}
        report = LAYOUT._object_scope_pages(
            pages,
            {"start_text": "（二）基础回归", "end_text": "（三）后续分析", "include_end_page": True},
        )
        self.assertTrue(report["all_pass"])
        self.assertEqual(report["pages"], [2, 3, 4])

    def test_table_locator_requires_complete_header_signature(self) -> None:
        markers = [{"table_index": 7, "caption": "", "header": ["链条", "记录", "二元组", "成年子女"]}]
        located, missing = LAYOUT._locate_table_pages(
            markers,
            {1: "链条 记录 unrelated", 2: "链条 记录 二元组 成年子女"},
        )
        self.assertEqual(located, {7: [2]})
        self.assertEqual(missing, [])

    def test_fixture_split_is_nine_left_two_right_contract(self) -> None:
        policy = LAYOUT.normalize_publication_layout_policy({})
        prose = ["Effect of X on Y through pathways other than M"] * 9
        controls = ["Counterfactual definition", "EYxm-Yxm*-Yx*m+Yx*m**Mx-Mx*"]
        self.assertEqual(sum(LAYOUT.classify_cell_text(value, policy) == "left" for value in prose), 9)
        self.assertEqual(sum(LAYOUT.classify_cell_text(value, policy) == "right" for value in controls), 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
