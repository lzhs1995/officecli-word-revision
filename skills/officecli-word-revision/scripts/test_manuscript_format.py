"""Regression fixtures for real inherited-format failures (no source changes)."""
import copy, tempfile, unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from manuscript_format import apply_contract,audit_contract,bool_value,get,object_registry,audit_native_clean

CONTRACT={
 'body':{'east_asia':'SimSun','latin':'Times New Roman','size_pt':10.5,'line_pt':18},
 'table':{'east_asia':'SimSun','latin':'Times New Roman','size_pt':10.5,'line_pt':14},
 'caption':{'east_asia':'KaiTi','latin':'Times New Roman','size_pt':10.5,'line_pt':18},
 'expected_objects':1}
PUB={'table':{'alignment_mode':'han-left-nonhan-right'}}

def fixture():
 d=Document();normal=d.styles['Normal'];pr=get(normal.element,'w:pPr')
 get(pr,'w:ind').set(qn('w:firstLineChars'),'200')
 inherited=d.styles.add_style('Inherited body',WD_STYLE_TYPE.PARAGRAPH);inherited.base_style=normal
 get(get(inherited.element,'w:rPr'),'w:sz').set(qn('w:val'),'24')
 d.add_paragraph('中文正文 Body text with inherited twelve point size',style=inherited)
 d.add_paragraph('表1 混合中西文格式测试')
 t=d.add_table(rows=3,cols=3)
 values=[['变量','Mean','95% CI'],['收入 income','1.250','1.096–1.421'],['','Long English text with more than three ordinary spaces','(0.1, 0.5)']]
 for row,vals in zip(t.rows,values):
  for c,v in zip(row.cells,vals):
   c.text=v
   for r in c.paragraphs[0].runs:get(get(r._r,'w:rPr'),'w:b').set(qn('w:val'),'0')
 return d

class FormatTests(unittest.TestCase):
 def check(self,d):
  with tempfile.TemporaryDirectory(prefix='wordrev-format-test-') as td:
   p=Path(td)/'fixture.docx';d.save(p);return audit_contract(p,CONTRACT,PUB)
 def good(self):
  d=fixture();apply_contract(d,CONTRACT,PUB);self.assertTrue(self.check(d)['all_pass']);return d
 def test_explicit_fonts_and_zero_indents(self):self.good()
 def test_opt_in_math_minimum_spacing(self):
  d=fixture();p=d.tables[0].cell(1,1).paragraphs[0]._p
  math=OxmlElement('m:oMath');p.append(math)
  contract=copy.deepcopy(CONTRACT);contract['table']['math_line_rule']='atLeast'
  apply_contract(d,contract,PUB)
  from journal_layout import audit_table_layout
  self.assertTrue(audit_table_layout(d,PUB,math_line_rule='atLeast')['all_pass'])
  self.assertFalse(audit_table_layout(d,PUB)['all_pass'])
  spacing=p.find('w:pPr/w:spacing',p.nsmap)
  self.assertEqual(spacing.get(qn('w:lineRule')),'atLeast')
  with tempfile.TemporaryDirectory(prefix='wordrev-math-test-')as td:
   path=Path(td)/'math.docx';d.save(path)
   self.assertTrue(audit_contract(path,contract,PUB)['all_pass'])
   spacing.set(qn('w:lineRule'),'exact');d.save(path)
   self.assertFalse(audit_contract(path,contract,PUB)['all_pass'])
 def test_false_is_false(self):
  for value in ['0','false','off']:
   e=OxmlElement('w:b');e.set(qn('w:val'),value);self.assertFalse(bool_value(e))
 def test_inherited_size_fails(self):
  d=self.good();r=d.paragraphs[0].runs[0]._r
  pr=r.find(qn('w:rPr'));pr.remove(pr.find(qn('w:sz')))
  self.assertFalse(self.check(d)['all_pass'])
 def test_wrong_chinese_font_fails(self):
  d=self.good();get(d.tables[0].cell(1,0).paragraphs[0].runs[0]._r.rPr,'w:rFonts').set(qn('w:eastAsia'),'Arial')
  self.assertFalse(self.check(d)['all_pass'])
 def test_wrong_western_font_fails(self):
  d=self.good();get(d.tables[0].cell(1,1).paragraphs[0].runs[0]._r.rPr,'w:rFonts').set(qn('w:hAnsi'),'Calibri')
  self.assertFalse(self.check(d)['all_pass'])
 def test_bold_fails(self):
  d=self.good();get(d.tables[0].cell(1,0).paragraphs[0].runs[0]._r.rPr,'w:b').set(qn('w:val'),'1')
  self.assertFalse(self.check(d)['all_pass'])
 def test_deleted_char_zero_fails(self):
  d=self.good();ind=d.tables[0].cell(1,0).paragraphs[0]._p.pPr.find(qn('w:ind'));del ind.attrib[qn('w:firstLineChars')]
  self.assertFalse(self.check(d)['all_pass'])
 def test_wrong_alignment_fails(self):
  d=self.good();get(d.tables[0].cell(1,0).paragraphs[0]._p.pPr,'w:jc').set(qn('w:val'),'center')
  self.assertFalse(self.check(d)['all_pass'])
 def test_word_omitted_default_left_alignment_passes(self):
  d=self.good();pr=d.tables[0].cell(1,0).paragraphs[0]._p.pPr
  pr.remove(pr.find(qn('w:jc')))
  self.assertTrue(self.check(d)['all_pass'])
 def test_inherited_wrong_alignment_still_fails(self):
  d=self.good();pr=d.tables[0].cell(1,0).paragraphs[0]._p.pPr
  pr.remove(pr.find(qn('w:jc')))
  get(get(d.styles['Normal'].element,'w:pPr'),'w:jc').set(qn('w:val'),'right')
  self.assertFalse(self.check(d)['all_pass'])
 def test_fully_deleted_prior_table_does_not_steal_caption(self):
  d=self.good();t=d.tables[0]._tbl;old=copy.deepcopy(t)
  for row in old.findall(qn('w:tr')):get(row,'w:trPr').append(OxmlElement('w:del'))
  t.addprevious(old)
  self.assertTrue(self.check(d)['all_pass'])
 def test_moved_caption_fails(self):
  d=self.good();cap=d.paragraphs[1]._p;cap.getparent().remove(cap);d.tables[0]._tbl.addnext(cap)
  self.assertFalse(self.check(d)['all_pass'])
 def test_body_bookmarks_do_not_break_caption_registry(self):
  d=self.good();table=d.tables[0]._tbl
  start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),'1');start.set(qn('w:name'),'TableAnchor')
  end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'1')
  table.addprevious(start);table.addnext(end)
  self.assertTrue(self.check(d)['all_pass'])
  self.assertEqual(object_registry(d,PUB)[0]['caption'],'表1 混合中西文格式测试')
 def test_missing_caption_fails(self):
  d=self.good();cap=d.paragraphs[1]._p;cap.getparent().remove(cap)
  self.assertFalse(self.check(d)['all_pass'])
 def test_native_changes(self):
  d=fixture();rows=apply_contract(d,CONTRACT,PUB,tracked=True)
  self.assertTrue(rows);self.assertTrue(d.element.xpath('.//w:rPrChange'));self.assertTrue(d.element.xpath('.//w:pPrChange'))
  self.assertTrue(self.check(d)['all_pass'])
  self.assertFalse(d.element.xpath('.//w:pPrChange/w:pPr/w:rPr'))
  self.assertFalse(d.element.xpath('.//w:rPrChange/w:rPr/w:ins|.//w:rPrChange/w:rPr/w:del'))
 def test_internal_border_fails(self):
  d=self.good();edge=get(get(d.tables[0].cell(1,1)._tc.tcPr,'w:tcBorders'),'w:bottom');edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'12')
  self.assertFalse(self.check(d)['all_pass'])
 def test_body_hanging_does_not_cancel_first_line(self):
  d=self.good();ind=d.paragraphs[0]._p.pPr.find(qn('w:ind'))
  self.assertEqual(ind.get(qn('w:firstLineChars')),'200');self.assertNotIn(qn('w:hangingChars'),ind.attrib)
 def test_last_paragraph_mark_is_not_a_clean_document(self):
  d=self.good();p=d.paragraphs[-1]._p;pr=get(get(p,'w:pPr'),'w:rPr');change=get(pr,'w:rPrChange');change.append(OxmlElement('w:rPr'))
  with tempfile.TemporaryDirectory() as td:
   path=Path(td)/'last-mark.docx';d.save(path)
   self.assertFalse(audit_native_clean(path)['all_pass'])
 def test_word_roundtrip_fixture_is_ready(self):
  # The independent owned Word save test is run by the project, not silently
  # launched by a unit test suite on machines without Microsoft Word.
  d=self.good();self.assertEqual(len(object_registry(d,PUB)),1)

if __name__=='__main__':unittest.main()
