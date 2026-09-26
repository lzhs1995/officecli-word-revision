from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from make_thesis_fixture import make_fixture, field, bookmark, element
from thesis_integrity import Audit, audit, digest, load_bound_contract, parse_fields, validate_contract, NS


class ThesisIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.path = self.root / "fixture.docx"
        self.spec = make_fixture(self.path)

    def mutate(self, action, part="word/document.xml"):
        with zipfile.ZipFile(self.path) as archive:
            data = {name: archive.read(name) for name in archive.namelist()}
        root = etree.fromstring(data[part])
        action(root)
        data[part] = etree.tostring(root)
        with zipfile.ZipFile(self.path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, payload in data.items():
                archive.writestr(name, payload)

    def failure(self, code):
        result = audit(self.path, self.spec)
        self.assertFalse(result["all_pass"], result)
        self.assertIn(code, [e["code"] for e in result["errors"]], result)
        return result

    def test_complete_synthetic_example_and_input_unchanged(self):
        before = digest(self.path)
        result = audit(self.path, self.spec)
        self.assertEqual(result["status"], "PASS", result)
        self.assertEqual(before, digest(self.path))
        self.assertEqual(result["observations"]["bibliography"]["groups"], {"en": 1, "zh": 1})
        self.assertEqual(result["observations"]["footnotes"]["count"], 2)
        self.assertEqual(len(result["observations"]["caption_pairs"]), 2)
        self.assertIn("word/footnotes.xml", result["observations"]["stories"])

    def test_empty_or_unknown_contract_rejected(self):
        from jsonschema import ValidationError
        for spec in ({"schema_version": "1.0", "blocks": []}, {**self.spec, "typo": True}):
            with self.assertRaises(ValidationError):
                validate_contract(spec)

    def test_bound_contract_fails_when_changed(self):
        path = self.root / "contract.json"
        path.write_text(json.dumps(self.spec))
        binding = {"path": str(path), "sha256": digest(path)}
        self.assertEqual(load_bound_contract(binding), self.spec)
        path.write_text(json.dumps({**self.spec, "caption_sequence_complete": False}))
        with self.assertRaisesRegex(ValueError, "QA_CONTRACT_CHANGED"):
            load_bound_contract(binding)

    def test_block_order_is_physical_not_contract_label(self):
        self.spec["blocks"][0], self.spec["blocks"][1] = self.spec["blocks"][1], self.spec["blocks"][0]
        self.failure("BLOCK_ORDER_OR_OVERLAP")

    def test_same_section_new_page_is_legal(self):
        self.assertEqual(audit(self.path, self.spec)["status"], "PASS")
        self.spec["blocks"][1]["start"] = "new_section"
        self.failure("BLOCK_SECTION_BOUNDARY")

    def test_false_pagebreak_and_inherited_pagebreak(self):
        def change(root):
            node = root.xpath('.//w:p[w:r/w:t="Abstract"]/w:pPr/w:pageBreakBefore', namespaces=NS)[0]
            node.set(qn("w:val"), "false")
        self.mutate(change)
        self.failure("BLOCK_PAGE_BOUNDARY")

    def test_roman_case_can_be_advisory_or_explicit(self):
        self.spec["sections"][0]["page_number"]["format"] = "lowerRoman"
        self.failure("PAGE_NUMBER_FORMAT")

    def test_accidental_appendix_restart_fails(self):
        self.mutate(lambda root: root.xpath('.//w:sectPr', namespaces=NS)[-1].find('w:pgNumType', NS).set(qn('w:start'), '1'))
        self.failure("PAGE_NUMBER_RESTART")

    def test_linked_header_is_not_independent_by_new_section(self):
        self.spec["sections"][-1]["stories"]["header_default"] = "independent"
        self.failure("STORY_LINK")

    def test_requested_section_missing_fails(self):
        self.spec["sections"].append({"section": 5, "start_type": "oddPage"})
        self.failure("SECTION_MISSING")

    def test_missing_ref_in_footer_fails(self):
        document = Document(self.path)
        field(document.sections[0].footer.add_paragraph(), ' REF DoesNotExist ', 'wrong')
        document.save(self.path)
        self.failure("REFERENCE_TARGET_MISSING_OR_EMPTY")

    def test_pageref_accepts_positional_bookmark(self):
        document = Document(self.path)
        p = document.add_paragraph()
        bookmark(p, "PositionOnly", 900)
        field(document.add_paragraph(), ' PAGEREF PositionOnly ', '10', simple=True)
        document.save(self.path)
        self.assertTrue(audit(self.path, self.spec)["all_pass"])

    def test_empty_ref_target_is_error(self):
        document = Document(self.path)
        bookmark(document.add_paragraph(), "EmptyNumber", 901)
        field(document.add_paragraph(), ' REF EmptyNumber ', '1')
        document.save(self.path)
        self.failure("REFERENCE_TARGET_MISSING_OR_EMPTY")

    def test_typed_contents_cannot_replace_toc(self):
        def change(root):
            n = next(n for n in root.iter(qn("w:instrText")) if '\\o "1-3"' in (n.text or ""))
            n.text = " QUOTE typed_contents "
        self.mutate(change)
        self.failure("TOC_COUNT")

    def test_native_style_based_figure_and_table_lists_are_supported_explicitly(self):
        def change(root):
            for n in root.iter(qn("w:instrText")):
                if '\\c "图"' in (n.text or ""):
                    n.text = ' TOC \\t "Figure Caption,4" \\h '
                elif '\\c "表"' in (n.text or ""):
                    n.text = ' TOC \\t "Table Caption,4" \\h '
        self.mutate(change)
        self.failure("TOC_COUNT")
        self.spec["fields"]["toc"].update(figure_styles=["Figure Caption"], table_styles=["Table Caption"])
        self.assertEqual(audit(self.path, self.spec)["status"], "PASS")
        self.spec["fields"]["toc"]["table_styles"].append("Figure Caption")
        self.failure("TOC_AMBIGUOUS_SELECTOR")

    def test_duplicate_bibliography_is_error(self):
        def change(root):
            n = next(n for n in root.iter(qn("w:instrText")) if "ZOTERO_BIBL" in (n.text or ""))
            p = n.getparent().getparent()
            last = p.getnext()
            first_copy, last_copy = copy.deepcopy(p), copy.deepcopy(last)
            last.addnext(first_copy)
            first_copy.addnext(last_copy)
        self.mutate(change)
        self.failure("ZOTERO_BIBLIOGRAPHY_COUNT")

    def test_malformed_bibliography_code_is_not_a_live_bibliography(self):
        def corrupt(root):
            n = next(n for n in root.iter(qn("w:instrText")) if "ZOTERO_BIBL" in (n.text or ""))
            n.text = " ADDIN ZOTERO_BIBL malformed CSL_BIBLIOGRAPHY "
        self.mutate(corrupt)
        self.failure("ZOTERO_BIBLIOGRAPHY_INVALID")

    def test_malformed_citation_item_produces_a_report_instead_of_crashing(self):
        def corrupt(root):
            n = next(n for n in root.iter(qn("w:instrText")) if "ZOTERO_ITEM" in (n.text or ""))
            n.text = ' ADDIN ZOTERO_ITEM CSL_CITATION {"citationItems":[42]} '
        self.mutate(corrupt)
        self.failure("ZOTERO_ITEM_INVALID")

    def test_caption_bound_to_wrong_neighbor_cannot_pass_position_only(self):
        document = Document(self.path)
        figure_caption = next(p for p in document.paragraphs
                              if p._p.xpath('.//w:bookmarkStart[@w:name="Figure11"]'))
        figure_caption.insert_paragraph_before("This narrative interrupts the picture-caption pair.")
        document.save(self.path)
        self.failure("CAPTION_PAIR_INTERRUPTED")

    def test_uri_repetition_between_fields_is_valid_within_field_is_not(self):
        def repeat(root):
            n = next(n for n in root.iter(qn("w:instrText")) if "ZOTERO_ITEM" in (n.text or ""))
            p = n.getparent().getparent()
            p.addnext(copy.deepcopy(p))
        self.mutate(repeat)
        self.assertTrue(audit(self.path, self.spec)["all_pass"])
        def corrupt(root):
            n = next(n for n in root.iter(qn("w:instrText")) if "ZOTERO_ITEM" in (n.text or ""))
            prefix, raw = n.text.split("{", 1)
            value = json.loads("{" + raw)
            value["citationItems"] *= 2
            n.text = prefix + json.dumps(value)
        self.mutate(corrupt)
        self.failure("ZOTERO_DUPLICATE_URI_WITHIN_FIELD")

    def test_dontupdate_is_not_accepted_as_live_refreshable(self):
        def corrupt(root):
            n = next(n for n in root.iter(qn("w:instrText")) if "ZOTERO_ITEM" in (n.text or ""))
            prefix, raw = n.text.split("{", 1)
            value = json.loads("{" + raw)
            value["citationItems"][0]["dontUpdate"] = True
            n.text = prefix + json.dumps(value)
        self.mutate(corrupt)
        self.failure("ZOTERO_DONT_UPDATE")

    def test_unclosed_field_in_note_cannot_be_repaired_by_other_story(self):
        def corrupt(root):
            root[-1][0][0].append(element("w:fldChar", fldCharType="begin"))
        self.mutate(corrupt, "word/footnotes.xml")
        self.failure("FIELD_UNCLOSED")

    def test_footnote_section_override_is_checked(self):
        def change(root):
            root.xpath('.//w:sectPr', namespaces=NS)[1].find('w:footnotePr/w:numRestart', NS).set(qn('w:val'), 'eachSect')
        self.mutate(change)
        self.failure("FOOTNOTE_SECTION_POLICY")

    def test_footnote_section_policy_survives_missing_document_defaults(self):
        self.mutate(lambda root: root.remove(root.find('w:footnotePr', NS)), "word/settings.xml")
        result = audit(self.path, self.spec)
        self.assertTrue(result["all_pass"], result)
        self.assertEqual({s["numRestart"] for s in result["observations"]["footnotes"]["effective_sections"]}, {"eachPage"})

    def test_lost_document_and_section_restart_is_not_continuous_pass(self):
        self.mutate(lambda root: root.remove(root.find('w:footnotePr', NS)), "word/settings.xml")
        def change(root):
            for pr in root.xpath('.//w:sectPr/w:footnotePr', namespaces=NS):
                pr.remove(pr.find('w:numRestart', NS))
        self.mutate(change)
        self.failure("FOOTNOTE_SECTION_POLICY")

    def test_footnote_omitted_position_defaults_to_page_bottom(self):
        self.mutate(lambda root: root.find('w:footnotePr', NS).remove(root.find('w:footnotePr/w:pos', NS)), "word/settings.xml")
        def change(root):
            for pr in root.xpath('.//w:sectPr/w:footnotePr', namespaces=NS):
                pr.remove(pr.find('w:pos', NS))
        self.mutate(change)
        self.assertTrue(audit(self.path, self.spec)["all_pass"])

    def test_orphan_footnote_is_not_silently_dropped(self):
        self.mutate(lambda root: root[-1].set(qn('w:id'), '3'), "word/footnotes.xml")
        self.failure("FOOTNOTE_REFERENCE_OR_ORPHAN")

    def test_baseline_is_not_superscript(self):
        def change(root):
            r = root.xpath('.//w:r[w:footnoteReference]', namespaces=NS)[0]
            r.find('w:rPr/w:vertAlign', NS).set(qn('w:val'), 'baseline')
        self.mutate(change)
        self.failure("FOOTNOTE_REFERENCE_NOT_SUPERSCRIPT")

    def test_explicit_object_anchor_prevents_caption_guessing(self):
        self.spec["captions"][0]["object"] = {"bookmark": "ObjectTable11"}
        self.failure("CAPTION_OBJECT_KIND")

    def test_caption_without_seq_is_not_live_number(self):
        def change(root):
            n = next(n for n in root.iter(qn("w:instrText")) if "SEQ 图" in (n.text or ""))
            n.text = ' QUOTE "1" '
        self.mutate(change)
        self.failure("CAPTION_SEQ")

    def test_figure_caption_keepnext_is_review_not_hard_failure(self):
        def change(root):
            p = root.xpath('.//w:p[w:bookmarkStart[@w:name="Figure11"]]', namespaces=NS)[0]
            p.find('w:pPr/w:keepNext', NS).set(qn('w:val'), 'true')
        self.mutate(change)
        result = audit(self.path, self.spec)
        self.assertEqual(result["status"], "REVIEW")
        self.assertTrue(result["all_pass"])

    def test_image_inherits_exact_body_line_height(self):
        def change(root):
            p = root.xpath('.//w:p[w:bookmarkStart[@w:name="ObjectFigure11"]]', namespaces=NS)[0]
            p.find('w:pPr', NS).remove(p.find('w:pPr/w:spacing', NS))
        self.mutate(change)
        self.failure("IMAGE_EXACT_LINE_HEIGHT")

    def test_caption_pair_missing_keepnext_fails(self):
        def change(root):
            p = root.xpath('.//w:p[w:bookmarkStart[@w:name="Table11"]]', namespaces=NS)[0]
            p.find('w:pPr/w:keepNext', NS).set(qn('w:val'), '0')
        self.mutate(change)
        self.failure("CAPTION_PAIR_NOT_KEPT")

    def test_character_indent_cannot_hide_behind_length_zero(self):
        self.mutate(lambda root: root.xpath('.//w:tbl//w:ind', namespaces=NS)[0].set(qn('w:firstLineChars'), '200'))
        self.failure("TABLE_INHERITED_INDENT")

    def test_table_row_cannot_clip_at_exact_height(self):
        self.mutate(lambda root: root.xpath('.//w:tbl//w:trHeight', namespaces=NS)[0].set(qn('w:hRule'), 'exact'))
        self.failure("TABLE_EXACT_ROW_HEIGHT")

    def test_false_bold_is_not_true_and_style_inheritance_is_checked(self):
        self.assertTrue(audit(self.path, self.spec)["all_pass"])
        self.mutate(lambda root: root.xpath('.//w:tbl//w:rPr/w:b', namespaces=NS)[0].set(qn('w:val'), 'true'))
        self.failure("TABLE_HEADER_BOLD")

    def test_bibliography_output_order_not_just_expected_count(self):
        def change(root):
            p = next(n for n in root.iter(qn("w:instrText")) if "ZOTERO_BIBL" in (n.text or "")).getparent()
            p.find('w:t', NS).text = '\n'.join(e['text'] for e in reversed(self.spec['bibliography']['entries']))
        self.mutate(change)
        self.failure("BIBLIOGRAPHY_VISIBLE_ORDER_OR_CONTENT")

    def test_language_is_declared_not_guessed_from_title_glyphs(self):
        entry = self.spec['bibliography']['entries'][0]
        old = entry['title']
        entry['title'] = '合成英文书名'
        entry['text'] = entry['text'].replace(old, entry['title'])
        self.mutate(lambda root: [setattr(n, 'text', n.text.replace(old, entry['title'])) for n in root.iter(qn('w:t')) if n.text and old in n.text])
        self.assertTrue(audit(self.path, self.spec)["all_pass"])

    def test_raw_newline_in_one_text_node_does_not_separate_bibliography_entries(self):
        def change(root):
            first = next(n for n in root.iter(qn('w:t')) if n.text == self.spec['bibliography']['entries'][0]['text'])
            second = next(n for n in root.iter(qn('w:t')) if n.text == self.spec['bibliography']['entries'][1]['text'])
            first.text += '\n' + second.text
            second.text = ''
        self.mutate(change)
        result = self.failure('BIBLIOGRAPHY_ENTRY_PARAGRAPHS')
        self.assertNotIn('BIBLIOGRAPHY_VISIBLE_ORDER_OR_CONTENT', {e['code'] for e in result['errors']})
        # Existing contracts can deliberately leave paragraph layout untested.
        self.spec['bibliography'].pop('entries_separate_paragraphs')
        self.assertTrue(audit(self.path, self.spec)['all_pass'])

    def test_full_author_key_breaks_same_initial_ties(self):
        entry = copy.deepcopy(self.spec['bibliography']['entries'][1])
        entry.update(id='zh-cai', author_display='蔡甲', author_sort='cai jia', text='蔡甲, 2021. 合成研究[M]. 示例出版社.')
        self.spec['bibliography']['entries'].append(entry)
        # Correct number of entries, wrong full-pinyin order.
        def change(root):
            n = next(n for n in root.iter(qn('w:t')) if n.text == self.spec['bibliography']['entries'][1]['text'])
            n.text += '\n' + entry['text']
        self.mutate(change)
        self.failure('BIBLIOGRAPHY_VISIBLE_ORDER_OR_CONTENT')


class FieldBoundaryTests(unittest.TestCase):
    def test_adjacent_split_and_nested_fields(self):
        d = Document()
        p = d.add_paragraph()
        field(p, ' REF A ', 'a')
        field(p, ' REF B ', 'b', simple=True)
        run = p.add_run()._r
        run.append(element('w:fldChar', fldCharType='begin'))
        for value in (' PA', 'GEREF A '):
            n = element('w:instrText'); n.text = value; run.append(n)
        run.append(element('w:fldChar', fldCharType='separate'))
        n = element('w:t'); n.text = '2'; run.append(n)
        run.append(element('w:fldChar', fldCharType='end'))
        fields, errors = parse_fields(d.element, 'word/document.xml')
        self.assertEqual(errors, [])
        self.assertEqual([f['instruction'].strip() for f in fields], ['REF A', 'REF B', 'PAGEREF A'])
        self.assertEqual([f['result'] for f in fields], ['a', 'b', '2'])
        # An outer IF result may contain a real nested REF; both retain values.
        outer = element('w:fldSimple', instr=' IF 1 = 1 "yes" "no" ')
        p._p.append(outer)
        inner = element('w:fldSimple', instr=' REF A ')
        r, t = element('w:r'), element('w:t'); t.text = 'a'; r.append(t); inner.append(r); outer.append(inner)
        fields, errors = parse_fields(d.element, 'word/document.xml')
        self.assertFalse(errors)
        self.assertEqual(fields[-1]['result'], 'a')


class PipelineGateTests(unittest.TestCase):
    setUp = ThesisIntegrityTests.setUp

    def test_pipeline_hash_binding_is_rechecked_before_cache(self):
        from word_revision_pipeline import Pipeline
        contract = self.root / 'contract.json'
        contract.write_text(json.dumps(self.spec))
        job = {'schema_version':'2.0','job_id':'fixture','scenario':'thesis_format', 'scope':'full_thesis',
               'source':str(self.path),'profile':'ruc-doctoral-2026','run_dir':str(self.root/'run'),
               'cache_dir':str(self.root/'cache'),'notebooklm':'disabled',
               'statistics':{'allow_execution':False,'consume_manifest_only':True},
               'output':{'tracked_formatting':False},
               'thesis_integrity':{'path':str(contract),'sha256':digest(contract)}}
        path = self.root/'job.json';path.write_text(json.dumps(job))
        pipeline = Pipeline(path)
        self.assertEqual(pipeline._input_hashes()['thesis_integrity'], digest(contract))
        self.assertIn(str(Path(__file__).with_name('thesis_integrity.py')), pipeline._implementation_hashes())
        contract.write_text(json.dumps({**self.spec, 'caption_sequence_complete':False}))
        with self.assertRaisesRegex(ValueError, 'QA_CONTRACT_CHANGED'):
            pipeline._input_hashes()

    def test_failed_integrity_is_a_final_adapter_failure(self):
        import thesis_format_adapter as adapter
        from types import SimpleNamespace
        self.spec['sections'][-1]['stories']['header_default'] = 'independent'
        contract = self.root/'contract.json';contract.write_text(json.dumps(self.spec))
        paths = {'accepted':str(self.path)}
        for key in ['semantic_before','semantic_after','refresh_log','metadata','adapter_qa_json','qa_md']:
            paths[key] = str(self.root/(key+'.json'))
        Path(paths['semantic_before']).write_text(json.dumps(adapter._xml_semantic_snapshot(self.path)))
        Path(paths['metadata']).write_text('{}')
        Path(paths['refresh_log']).write_text('{"errors":0}')
        job = {'profile':'ruc-doctoral-2026','source':str(self.path),'scope':'full_thesis',
               'thesis_integrity':{'path':str(contract),'sha256':digest(contract)}}
        _, profile = adapter.load_profile(job)
        response = SimpleNamespace(stdout=json.dumps({'success':True,'results':[{'format':profile['document_props']}]}))
        with patch.object(adapter, 'office', return_value=response):
            with self.assertRaisesRegex(RuntimeError, 'thesis_integrity_failed'):
                adapter.verify(job, paths, {}, 0)
        result = json.loads(Path(paths['adapter_qa_json']).read_text())
        self.assertFalse(result['all_pass'])
        self.assertIn('STORY_LINK', [e['code'] for e in result['thesis_integrity']['errors']])


class SharedLockTests(unittest.TestCase):
    def test_adapter_runner_nested_same_thread_but_other_thread_excluded(self):
        from thesis_format_adapter import _word_process_lock
        from word_runtime import word_lock
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'word.lock'
            def contender():
                with word_lock('other-thread', timeout=.15, path=path):
                    return 'incorrectly acquired'
            with _word_process_lock(path, 1):
                with word_lock('nested-runner', timeout=.1, path=path):
                    with ThreadPoolExecutor(max_workers=1) as pool:
                        with self.assertRaises(TimeoutError):
                            pool.submit(contender).result(timeout=2)
            with word_lock('after-return', timeout=.1, path=path):
                self.assertTrue(path.exists())


if __name__ == '__main__':
    unittest.main()
