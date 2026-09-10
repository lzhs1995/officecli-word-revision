"""Adversarial, Word-free contracts for thesis adapter hardening."""

from __future__ import annotations

import importlib.util
import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import unittest
import zlib
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from lxml import etree


ADAPTER_PATH = Path(__file__).parent / "thesis_format_adapter.py"
SPEC = importlib.util.spec_from_file_location("thesis_adapter", ADAPTER_PATH)
if SPEC is None or SPEC.loader is None:
    raise ImportError(f"Cannot load {ADAPTER_PATH}")
adapter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(adapter)


def _png_bytes(red: int, green: int, blue: int) -> bytes:
    def chunk(kind: bytes, payload: bytes) -> bytes:
        checksum = zlib.crc32(kind + payload) & 0xFFFFFFFF
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", checksum)

    signature = b"\x89PNG\r\n\x1a\n"
    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = b"\x00" + bytes((red, green, blue))
    return signature + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(pixels)) + chunk(b"IEND", b"")


def _rewrite_members(path: Path, replacements: dict[str, bytes]) -> None:
    temporary = path.with_name(f".{path.name}.rewrite")
    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    temporary.replace(path)


def _document_with_images(path: Path, colors: list[tuple[int, int, int]]) -> None:
    document = Document()
    document.add_paragraph("same text")
    for index, color in enumerate(colors):
        image_path = path.parent / f"image-{index}.png"
        image_path.write_bytes(_png_bytes(*color))
        document.add_picture(str(image_path))
    document.save(path)


def _inject_equation(path: Path, text: str, attributes: list[tuple[str, str]] | None = None) -> None:
    with ZipFile(path, "r") as package:
        document_xml = package.read("word/document.xml")
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    }
    root = etree.fromstring(document_xml)
    body = root.find(".//w:body", namespaces=namespaces)
    paragraph = etree.SubElement(body, f"{{{namespaces['w']}}}p")
    equation = etree.SubElement(paragraph, f"{{{namespaces['m']}}}oMath")
    for name, value in attributes or []:
        equation.set(name, value)
    equation.text = text
    _rewrite_members(
        path,
        {"word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8")},
    )


class MediaSemanticTest(unittest.TestCase):
    def test_different_image_bytes_change_payload_snapshot(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            first = root / "first.docx"
            second = root / "second.docx"
            _document_with_images(first, [(255, 0, 0)])
            shutil.copy2(first, second)
            with ZipFile(second) as package:
                media_name = next(name for name in package.namelist() if name.startswith("word/media/"))
            _rewrite_members(second, {media_name: _png_bytes(0, 0, 255)})
            self.assertNotEqual(
                adapter._xml_semantic_snapshot(first)["media_payload_sha256"],
                adapter._xml_semantic_snapshot(second)["media_payload_sha256"],
            )

    def test_swapped_image_relationships_change_usage_not_payload_bag(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            first = root / "first.docx"
            second = root / "second.docx"
            _document_with_images(first, [(255, 0, 0), (0, 0, 255)])
            shutil.copy2(first, second)
            rels_name = "word/_rels/document.xml.rels"
            with ZipFile(second) as package:
                rels_root = etree.fromstring(package.read(rels_name))
            relationship_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
            image_relationships = [
                node
                for node in rels_root.findall(f"{{{relationship_ns}}}Relationship")
                if (node.get("Type") or "").endswith("/image")
            ]
            self.assertGreaterEqual(len(image_relationships), 2)
            targets = [image_relationships[0].get("Target"), image_relationships[1].get("Target")]
            image_relationships[0].set("Target", targets[1])
            image_relationships[1].set("Target", targets[0])
            _rewrite_members(second, {rels_name: etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8")})
            before = adapter._xml_semantic_snapshot(first)
            after = adapter._xml_semantic_snapshot(second)
            self.assertEqual(before["media_payload_sha256"], after["media_payload_sha256"])
            self.assertNotEqual(before["media_usage_sha256"], after["media_usage_sha256"])


class EquationSemanticTest(unittest.TestCase):
    def test_different_equation_content_changes_snapshot(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            snapshots = []
            for index, text in enumerate(("x+1", "y+2")):
                path = root / f"equation-{index}.docx"
                Document().save(path)
                _inject_equation(path, text)
                snapshots.append(adapter._xml_semantic_snapshot(path))
            self.assertNotEqual(snapshots[0]["equation_xml_sha256"], snapshots[1]["equation_xml_sha256"])

    def test_attribute_order_is_canonicalized(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            paths = [root / "first.docx", root / "second.docx"]
            namespace = "http://schemas.openxmlformats.org/officeDocument/2006/math"
            for path in paths:
                Document().save(path)
            _inject_equation(paths[0], "x", [(f"{{{namespace}}}b", "2"), (f"{{{namespace}}}a", "1")])
            _inject_equation(paths[1], "x", [(f"{{{namespace}}}a", "1"), (f"{{{namespace}}}b", "2")])
            self.assertEqual(
                adapter._xml_semantic_snapshot(paths[0])["equation_xml_sha256"],
                adapter._xml_semantic_snapshot(paths[1])["equation_xml_sha256"],
            )


class FootnoteNormalizationTest(unittest.TestCase):
    def test_footnote_text_is_invariant_to_run_splits(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            paths = [root / "first.docx", root / "second.docx"]
            for path in paths:
                document = Document()
                document.add_paragraph("body")
                document.save(path)
            namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            single = f'''<w:footnotes xmlns:w="{namespace}"><w:footnote w:id="1"><w:p><w:r><w:t>Identical footnote text</w:t></w:r></w:p></w:footnote></w:footnotes>'''
            split = f'''<w:footnotes xmlns:w="{namespace}"><w:footnote w:id="1"><w:p><w:r><w:t>Identical </w:t></w:r><w:r><w:t>footnote </w:t></w:r><w:r><w:t>text</w:t></w:r></w:p></w:footnote></w:footnotes>'''
            _rewrite_members(paths[0], {"word/footnotes.xml": single.encode()})
            _rewrite_members(paths[1], {"word/footnotes.xml": split.encode()})
            self.assertEqual(
                adapter._xml_semantic_snapshot(paths[0])["footnote_text_sha256"],
                adapter._xml_semantic_snapshot(paths[1])["footnote_text_sha256"],
            )


class StructureAndAuditTest(unittest.TestCase):
    def _structure_status(self, scope: str) -> tuple[str, str]:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        source = root / "document.docx"
        document = Document()
        document.add_paragraph("Only chapter content")
        document.save(source)
        result = adapter.augment_audit(
            {"profile": "ruc-doctoral-2026", "scope": scope, "structure_policy": "validate_only"},
            source,
            root,
            {},
        )
        rows = json.loads(Path(result["format_diff"]).read_text(encoding="utf-8"))
        row = next(item for item in rows if item["area"] == "structure")
        return row["status"], row["current"]

    def test_full_thesis_missing_blocks_are_review(self):
        status, current = self._structure_status("full_thesis")
        self.assertEqual(status, "review")
        self.assertIn("中文摘要", current)

    def test_chapter_missing_full_thesis_blocks_is_pass(self):
        status, _ = self._structure_status("chapter")
        self.assertEqual(status, "pass")

    def test_audit_uses_zero_value_format_equivalence(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "document.docx"
            Document().save(source)
            results = {
                "document": {
                    "success": True,
                    "data": {"results": [{"format": {"docDefaults.spaceBefore": "0"}}]},
                }
            }
            result = adapter.augment_audit(
                {"profile": "ruc-doctoral-2026", "scope": "chapter"}, source, root, results
            )
            rows = json.loads(Path(result["format_diff"]).read_text(encoding="utf-8"))
            row = next(item for item in rows if item["target"] == "docDefaults.spaceBefore")
            self.assertEqual(row["status"], "pass")


class StyleSelectorValidationTest(unittest.TestCase):
    def test_nonexistent_style_selector_raises_from_real_source_styles(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "chapter.docx"
            Document().save(source)
            _, profile = adapter.load_profile({"profile": "ruc-doctoral-2026"})
            with self.assertRaisesRegex(ValueError, "DefinitelyMissing"):
                adapter._require_style_map(
                    {
                        "scope": "chapter",
                        "style_map": {"body": "/styles/DefinitelyMissing"},
                        "required_style_roles": ["body"],
                    },
                    profile,
                    source,
                )

    def test_existing_style_selector_passes_from_real_source_styles(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "chapter.docx"
            document = Document()
            document.add_paragraph("text")
            document.save(source)
            _, profile = adapter.load_profile({"profile": "ruc-doctoral-2026"})
            adapter._require_style_map(
                {
                    "scope": "chapter",
                    "style_map": {"body": "/styles/Normal"},
                    "required_style_roles": ["body"],
                },
                profile,
                source,
            )


class ProfileGeometryTest(unittest.TestCase):
    def test_all_geometry_values_use_canonical_profile_keys(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "document.docx"
            Document().save(path)
            profile = {
                "document_props": {
                    "pageWidth": "19cm",
                    "pageHeight": "27cm",
                    "marginTop": "1.1cm",
                    "marginBottom": "1.2cm",
                    "marginLeft": "1.3cm",
                    "marginRight": "1.4cm",
                    "marginHeader": "0.7cm",
                    "marginFooter": "0.8cm",
                },
                "style_roles": {"body": {}},
                "header": {"font_east_asia": "宋体", "font_latin": "Times New Roman", "size_pt": 10.5},
                "footer": {"font_latin": "Times New Roman", "size_pt": 10.5},
                "table": {"font_east_asia": "宋体", "font_latin": "Times New Roman", "size_pt": 10.5},
                "page_numbering": {"front_matter": "lowerRoman", "body": "decimal", "body_start": 1},
            }
            adapter._apply_structural_formatting(path, {"scope": "chapter"}, profile)
            section = Document(path).sections[0]
            expected = {
                "page_width": 19.0,
                "page_height": 27.0,
                "top_margin": 1.1,
                "bottom_margin": 1.2,
                "left_margin": 1.3,
                "right_margin": 1.4,
                "header_distance": 0.7,
                "footer_distance": 0.8,
            }
            for attribute, centimeters in expected.items():
                self.assertAlmostEqual(getattr(section, attribute).cm, centimeters, places=2, msg=attribute)


HOLDER_CODE = r'''
import importlib.util
import os
import sys
import time
from pathlib import Path

spec = importlib.util.spec_from_file_location("lock_adapter", Path(sys.argv[1]))
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
with module._word_process_lock(Path(sys.argv[2]), 5.0):
    Path(sys.argv[3]).write_text(str(os.getpid()), encoding="utf-8")
    time.sleep(float(sys.argv[4]))
'''


class WordProcessLockTest(unittest.TestCase):
    def _start_holder(self, root: Path, duration: float) -> tuple[subprocess.Popen, Path, Path]:
        lock_path = root / "word.lock"
        ready = root / "ready.txt"
        environment = dict(os.environ)
        environment["PYTHONDONTWRITEBYTECODE"] = "1"
        process = subprocess.Popen(
            [sys.executable, "-c", HOLDER_CODE, str(ADAPTER_PATH), str(lock_path), str(ready), str(duration)],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            env=environment,
        )
        deadline = time.monotonic() + 10
        while not ready.exists() and process.poll() is None and time.monotonic() < deadline:
            time.sleep(0.05)
        if not ready.exists():
            process.wait(timeout=5)
            self.fail(f"holder did not acquire lock: returncode={process.returncode}")
        return process, lock_path, ready

    def test_second_process_waits_and_lock_inode_is_persistent(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            process, lock_path, _ = self._start_holder(root, 0.8)
            inode_while_held = lock_path.stat().st_ino
            started = time.monotonic()
            with adapter._word_process_lock(lock_path, 3.0):
                elapsed = time.monotonic() - started
            process.wait(timeout=5)
            self.assertGreater(elapsed, 0.5)
            self.assertTrue(lock_path.exists())
            self.assertEqual(inode_while_held, lock_path.stat().st_ino)
            with adapter._word_process_lock(lock_path, 1.0):
                self.assertEqual(inode_while_held, lock_path.stat().st_ino)

    def test_timeout_names_actual_owner_path_and_wait(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            process, lock_path, ready = self._start_holder(root, 0.8)
            owner_pid = ready.read_text(encoding="utf-8")
            with self.assertRaises(TimeoutError) as caught:
                with adapter._word_process_lock(lock_path, 0.2):
                    pass
            message = str(caught.exception)
            self.assertIn(owner_pid, message)
            self.assertIn(str(lock_path), message)
            self.assertIn("after", message)
            process.wait(timeout=5)


class WordAutomationSafetyTest(unittest.TestCase):
    def test_actual_scripts_restore_alerts_and_close_only_owned_refs(self):
        compare_script = adapter._word_compare_script()
        refresh_script = adapter._word_refresh_script()
        pdf_script = adapter._word_pdf_script()
        for script in (compare_script, refresh_script, pdf_script):
            self.assertIn("on error errMsg number errNum", script)
            self.assertIn("set display alerts to previousAlerts", script)
            self.assertNotIn("document 1", script)
        for reference in ("originalDoc", "revisedDoc", "comparedDoc"):
            self.assertIn(f"if {reference} is not missing value then close {reference} saving no", compare_script)
        self.assertIn("if docRef is not missing value then close docRef saving no", refresh_script)
        self.assertIn("if docRef is not missing value then close docRef saving no", pdf_script)

    def test_refresh_production_entry_point_acquires_lock(self):
        with patch.object(adapter, "_word_process_lock") as lock, patch.object(
            adapter, "_word_refresh_selected_unlocked", return_value={"updated": 1}
        ) as unlocked:
            result = adapter._word_refresh_selected(Path("example.docx"))
        self.assertEqual(result, {"updated": 1})
        lock.assert_called_once_with(adapter._word_lock_path(), adapter.WORD_LOCK_TIMEOUT_SECONDS)
        unlocked.assert_called_once_with(Path("example.docx"))

    def test_compare_production_entry_point_acquires_lock(self):
        with patch.object(adapter, "_word_process_lock") as lock, patch.object(
            adapter, "_word_compare_formatting_unlocked", return_value={"revisions": 1}
        ) as unlocked:
            result = adapter._word_compare_formatting(
                Path("original.docx"), Path("revised.docx"), Path("tracked.docx"), "Codex"
            )
        self.assertEqual(result, {"revisions": 1})
        lock.assert_called_once_with(adapter._word_lock_path(), adapter.WORD_LOCK_TIMEOUT_SECONDS)
        unlocked.assert_called_once()

    def test_pdf_production_entry_point_acquires_lock(self):
        with patch.object(adapter, "_word_process_lock") as lock, patch.object(
            adapter, "_export_pdf_unlocked"
        ) as unlocked:
            adapter.export_pdf(Path("source.docx"), Path("output.pdf"), print_markup=False)
        lock.assert_called_once_with(adapter._word_lock_path(), adapter.WORD_LOCK_TIMEOUT_SECONDS)
        unlocked.assert_called_once_with(
            Path("source.docx"), Path("output.pdf"), print_markup=False
        )


class FinalSemanticGateTest(unittest.TestCase):
    def test_new_payload_invariants_are_mandatory(self):
        required = {"media_payload_sha256", "media_usage_sha256", "equation_xml_sha256"}
        self.assertTrue(required.issubset(set(adapter.SEMANTIC_INVARIANT_KEYS)))


if __name__ == "__main__":
    unittest.main(verbosity=2)
