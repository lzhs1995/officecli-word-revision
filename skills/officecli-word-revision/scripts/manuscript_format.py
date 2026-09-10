"""Explicit manuscript typography, native format revisions and final-file QA.

No source defaults are changed. Font policies are resolved by the job; this
module does not prescribe a school's fonts. Diagram carrier tables are figures.
"""
from __future__ import annotations
import copy, hashlib, json, re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from journal_layout import (_table_records, _table_is_diagram,
    normalize_publication_layout_policy, _column_semantics, _empty_cell_alignment,
    classify_cell_text, _set_cell_margin, _set_paragraph_format,
    canonicalize_property_order)

CAP = re.compile(r'^(图|表|Figure|Table)\s*([0-9A-Z]+(?:\.[0-9]+)*[A-Z]?)\s+\S',re.I)
NOTE = re.compile(r'^(注[：:]|Note[：:]|来源[：:])',re.I)
HEAD = re.compile(r'^([一二三四五六七八九十]+[、.]|[（(][一二三四五六七八九十]+[）)])')
FALSE={'0','false','off'}

def bool_value(e):
    return e is not None and e.get(qn('w:val'),'true').lower() not in FALSE

def audit_native_clean(path):
    """Include last paragraph-mark changes omitted by some UI enumerators."""
    d=Document(path);types=Counter()
    names={'ins','del','moveFrom','moveTo','rPrChange','pPrChange','tblPrChange',
           'tcPrChange','trPrChange','sectPrChange','numberingChange','cellIns','cellDel','cellMerge'}
    for part in d.part.package.parts:
        if not hasattr(part,'element'):continue
        for e in part.element.iter():
            name=e.tag.split('}')[-1]
            if name in names:types[name]+=1
    return {'all_pass':not types,'native_revision_count':sum(types.values()),'types':dict(types),
            'scope':'all parsed Word parts, including final paragraph mark'}

def visible(e):
    for a in [e]+list(e.iterancestors()):
        if a.tag in {qn('w:del'),qn('w:moveFrom')}:return False
        if a.tag==qn('w:tr') and a.find('w:trPr/w:del',a.nsmap) is not None:return False
    return True

def text(e):
    return ''.join(n.text or '' for n in e.iterdescendants(qn('w:t')) if visible(n))

def get(parent,tag):
    e=parent.find(qn(tag))
    if e is None:e=OxmlElement(tag);parent.append(e)
    return e

def effective_run(document,p,r):
    """Resolve ordinary run typography across defaults and both style chains.

    The formatter writes explicit properties, so conditional table styles may
    not override this contract. Missing explicit run properties are QA failures.
    """
    vals={'b':False,'bCs':False}
    def apply(pr):
        if pr is None:return
        f=pr.find(qn('w:rFonts'))
        if f is not None:
            for key in ['ascii','hAnsi','eastAsia','cs']:
                if f.get(qn('w:'+key)) is not None:vals[key]=f.get(qn('w:'+key))
        for key in ['sz','szCs','b','bCs']:
            n=pr.find(qn('w:'+key))
            if n is not None:vals[key]=bool_value(n) if key in ['b','bCs'] else n.get(qn('w:val'))
    styles={s.style_id:s.element for s in document.styles}
    def chain(sid,seen=None):
        seen=set() if seen is None else seen
        if not sid or sid in seen or sid not in styles:return
        seen.add(sid);s=styles[sid];b=s.find(qn('w:basedOn'))
        if b is not None:chain(b.get(qn('w:val')),seen)
        apply(s.find(qn('w:rPr')))
    defaults=document.styles.element.find('w:docDefaults/w:rPrDefault/w:rPr',document.styles.element.nsmap)
    apply(defaults)
    ps=p.find('w:pPr/w:pStyle',p.nsmap)
    default=next((s.style_id for s in document.styles if s.element.get(qn('w:default'))=='1' and s.element.get(qn('w:type'))=='paragraph'),None)
    chain(ps.get(qn('w:val')) if ps is not None else default)
    rs=r.find('w:rPr/w:rStyle',r.nsmap)
    if rs is not None:chain(rs.get(qn('w:val')))
    apply(r.find(qn('w:rPr')))
    return vals

def effective_paragraph(document,p):
    """Word removes redundant direct properties on save; resolve inheritance."""
    values={};styles={s.style_id:s.element for s in document.styles}
    def apply(pr):
        if pr is None:return
        for child in pr:
            name=child.tag.split('}')[-1]
            if name in ['rPr','pPrChange','sectPr']:continue
            values.setdefault(name,{}).update({k.split('}')[-1]:v for k,v in child.attrib.items()})
    def chain(sid,seen=None):
        seen=set() if seen is None else seen
        if not sid or sid in seen or sid not in styles:return
        seen.add(sid);s=styles[sid];b=s.find(qn('w:basedOn'))
        if b is not None:chain(b.get(qn('w:val')),seen)
        apply(s.find(qn('w:pPr')))
    apply(document.styles.element.find('w:docDefaults/w:pPrDefault/w:pPr',document.styles.element.nsmap))
    default=next((s.style_id for s in document.styles if s.element.get(qn('w:default'))=='1' and s.element.get(qn('w:type'))=='paragraph'),None)
    chain(style_id(p) or default);apply(p.find(qn('w:pPr')))
    return values

def set_fonts(pr,policy,bold=None):
    f=get(pr,'w:rFonts');f.attrib.clear()
    for k in ['ascii','hAnsi','cs']:f.set(qn('w:'+k),policy['latin'])
    f.set(qn('w:eastAsia'),policy['east_asia'])
    for k in ['sz','szCs']:get(pr,'w:'+k).set(qn('w:val'),str(round(policy['size_pt']*2)))
    if bold is not None:
        for k in ['b','bCs']:get(pr,'w:'+k).set(qn('w:val'),'1' if bold else '0')

def style_id(p):
    e=p.find('w:pPr/w:pStyle',p.nsmap)
    return e.get(qn('w:val')) if e is not None else None

def paragraph_role(d,p):
    t=text(p).strip()
    if p.xpath('.//w:drawing|.//w:pict'):return 'image'
    if CAP.match(t):return 'caption'
    if NOTE.match(t):return 'note'
    sid=style_id(p)
    if HEAD.match(t):return 'heading'
    if sid:
        try:
            s=d.styles[sid]
            if any(k in s.name for k in ['一级','二级','三级','四级','Heading','Title']):return 'heading'
        except KeyError:pass
    if p.xpath('.//m:oMath') and not t:return 'equation'
    return 'body'

def table_kind(t,policy):
    # Use the accepted-view rows when a cumulative redline still has old rows.
    tmp=copy.deepcopy(t._tbl)
    for r in list(tmp.findall(qn('w:tr'))):
        if r.find('w:trPr/w:del',r.nsmap) is not None:tmp.remove(r)
    if not tmp.findall(qn('w:tr')):return 'deleted'
    from docx.table import Table
    return 'figure' if _table_is_diagram(_table_records(Table(tmp,t._parent)),policy) else 'table'

class Recorder:
    """Store the old properties once; never nest an existing change record."""
    TYPES={'pPr':'pPrChange','rPr':'rPrChange','tcPr':'tcPrChange','trPr':'trPrChange','tblPr':'tblPrChange'}
    def __init__(self,d,author,date,enabled):
        self.author,self.date,self.enabled=author,date,enabled;self.rows=[]
        ids=[int(e.get(qn('w:id'))) for e in d.element.iter() if e.get(qn('w:id'),'').isdigit()]
        self.next=max(ids+[1000000])+1
    def edit(self,parent,tag,fn,location):
        name=tag.split(':')[-1];old=parent.find(qn(tag));previous=copy.deepcopy(old)
        if old is None:
            old=OxmlElement(tag);parent.insert(0,old)
        before=hashlib.sha256(bytes(str(old.xml),'utf8')).hexdigest()
        fn(old)
        after=hashlib.sha256(bytes(str(old.xml),'utf8')).hexdigest()
        if before==after:return
        if self.enabled:
            if name=='pPr':
                # ParagraphPropertyChange contains PPrBase, not the paragraph
                # mark's rPr. Track the mark separately to avoid resurrecting
                # old ins/del markers when all revisions are rejected.
                mark=old.find(qn('w:rPr'))
                prior_mark=previous.find(qn('w:rPr')) if previous is not None else None
                if mark is not None and mark.find(qn('w:rPrChange')) is None:
                    original=copy.deepcopy(prior_mark) if prior_mark is not None else OxmlElement('w:rPr')
                    for child in list(original):
                        if child.tag in {qn('w:'+k) for k in ['ins','del','moveFrom','moveTo','rPrChange']}:original.remove(child)
                    rc=OxmlElement('w:rPrChange');rc.set(qn('w:id'),str(self.next));self.next+=1
                    rc.set(qn('w:author'),self.author);rc.set(qn('w:date'),self.date);rc.append(original);mark.append(rc)
            change=old.find(qn('w:'+self.TYPES[name]))
            if change is None:
                change=OxmlElement('w:'+self.TYPES[name]);change.set(qn('w:id'),str(self.next));self.next+=1
                change.set(qn('w:author'),self.author);change.set(qn('w:date'),self.date)
                prior=previous if previous is not None else OxmlElement(tag)
                excluded=['ins','del','moveFrom','moveTo',self.TYPES[name]]
                if name=='pPr':excluded+=['rPr','sectPr']
                for child in list(prior):
                    if child.tag in {qn('w:'+k) for k in excluded}:prior.remove(child)
                change.append(prior);old.append(change)
        self.rows.append({'location':location,'property':name,'old_hash':before,'new_hash':after})

def _pformat(pr,role,policy):
    ind=get(pr,'w:ind')
    for key in ['firstLine','hanging','left','right','start','end','firstLineChars','hangingChars','leftChars','rightChars','startChars','endChars']:
        ind.set(qn('w:'+key),'0')
    if role=='body':
        ind.set(qn('w:firstLineChars'),'200');ind.set(qn('w:firstLine'),str(round(policy['size_pt']*40)))
        # A zero hanging attribute can still take precedence over firstLine.
        for key in ['hanging','hangingChars']:ind.attrib.pop(qn('w:'+key),None)
        get(pr,'w:keepNext').set(qn('w:val'),'0')
    sp=get(pr,'w:spacing');sp.attrib.clear()
    sp.set(qn('w:line'),str(round(policy['line_pt']*20)));sp.set(qn('w:lineRule'),'exact')
    for key in ['before','after','beforeLines','afterLines']:sp.set(qn('w:'+key),'0')
    sp.set(qn('w:beforeAutospacing'),'0');sp.set(qn('w:afterAutospacing'),'0')
    get(pr,'w:jc').set(qn('w:val'),'center' if role=='caption' else 'left' if role=='note' else 'both')
    mark=get(pr,'w:rPr');set_fonts(mark,policy,False if role in ['caption','note'] else None)

def apply_contract(d,contract,publication,*,tracked=False,author='Codex',date='2026-09-07T00:00:00Z'):
    policy=normalize_publication_layout_policy(publication);rec=Recorder(d,author,date,tracked)
    for i,p in enumerate(d.element.body.findall(qn('w:p'))):
        if not visible(p) or not text(p).strip():continue
        role=paragraph_role(d,p)
        if role not in ['body','caption','note']:continue
        fmt=contract['caption' if role=='caption' else 'body']
        rec.edit(p,'w:pPr',lambda pr:_pformat(pr,role,fmt),f'p:{i}:{role}')
        for j,r in enumerate(p.xpath('.//w:r[w:t]')):
            if visible(r):rec.edit(r,'w:rPr',lambda pr:set_fonts(pr,fmt,False if role=='caption' else None),f'p:{i}:r:{j}')
    for ti,t in enumerate(d.tables):
        if table_kind(t,policy)!='table':continue
        records=[r for r in _table_records(t) if visible(r['tc'])];sem=_column_semantics(records,policy)
        fmt=contract['table'];visible_rows=[r for r in t._tbl.findall(qn('w:tr')) if visible(r)]
        first_row=min(r['row'] for r in records);last_row=max(r['row'] for r in records)
        header_last=first_row+1 if '变量特定有效N' in text(t._tbl) else first_row
        def tableprops(pr):
            borders=get(pr,'w:tblBorders');borders.clear()
            for side in ['top','left','bottom','right','insideH','insideV']:
                b=get(borders,'w:'+side);b.set(qn('w:val'),'single' if side in ['top','bottom'] else 'nil')
                if side in ['top','bottom']:b.set(qn('w:sz'),'12');b.set(qn('w:color'),'000000')
        rec.edit(t._tbl,'w:tblPr',tableprops,f't:{ti}:borders')
        for ri,row in enumerate(visible_rows):
            def rowprops(pr):
                get(pr,'w:cantSplit')
                if ri<=header_last-first_row:get(pr,'w:tblHeader')
            rec.edit(row,'w:trPr',rowprops,f't:{ti}:row:{ri}')
        for record in records:
            tc=record['tc'];cls=classify_cell_text(text(tc),policy)
            align=_empty_cell_alignment(record,sem) if cls=='empty' else cls
            for pi,p in enumerate(tc.findall(qn('w:p'))):
                if not visible(p):continue
                loc=f't:{ti}:r:{record["row"]}:c:{record["cell"]}:p:{pi}'
                def pprops(pr):
                    _pformat(pr,'note',fmt);get(pr,'w:jc').set(qn('w:val'),align)
                    if p.xpath('.//m:oMath') and fmt.get('math_line_rule')=='atLeast':
                        get(pr,'w:spacing').set(qn('w:lineRule'),'atLeast')
                rec.edit(p,'w:pPr',pprops,loc)
                for j,r in enumerate(p.xpath('.//w:r[w:t]')):
                    if visible(r):rec.edit(r,'w:rPr',lambda pr:set_fonts(pr,fmt,False),loc+f':run:{j}')
            def tcprops(pr):
                borders=get(pr,'w:tcBorders');borders.clear()
                for side in ['top','left','bottom','right','insideH','insideV']:
                    b=get(borders,'w:'+side);b.set(qn('w:val'),'nil')
                if record['row']==first_row:
                    b=get(borders,'w:top');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'12');b.set(qn('w:color'),'000000')
                if record['row']==header_last:
                    b=get(borders,'w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'000000')
                if record['row']==last_row:
                    b=get(borders,'w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'12');b.set(qn('w:color'),'000000')
                mar=get(pr,'w:tcMar')
                for side in ['left','right']:
                    m=get(mar,'w:'+side);m.set(qn('w:type'),'dxa');m.set(qn('w:w'),'102')
            rec.edit(tc,'w:tcPr',tcprops,f't:{ti}:c:{record["row"]}:{record["cell"]}')
    canonicalize_property_order(d)
    return rec.rows

def object_registry(d,publication):
    policy=normalize_publication_layout_policy(publication);els=list(d.element.body);tables={id(t._tbl):t for t in d.tables};result=[]
    for i,e in enumerate(els):
        if not visible(e):continue
        if e.tag==qn('w:tbl'):
            kind=table_kind(tables[id(e)],policy)
            if kind=='deleted':continue
        elif any(visible(x) for x in e.iterdescendants(qn('w:drawing'),qn('w:pict'))) and not any(e.iterdescendants(qn('m:oMath'))):kind='figure'
        else:continue
        step=-1 if kind=='table' else 1;cap=None;cap_index=None
        for j in range(i+step,len(els) if step==1 else -1,step):
            x=els[j]
            if not visible(x):continue
            if x.tag==qn('w:tbl') and table_kind(tables[id(x)],policy)=='deleted':continue
            if x.tag==qn('w:tbl') or any(visible(y) for y in x.iterdescendants(qn('w:drawing'),qn('w:pict'))):break
            s=text(x).strip()
            if not s:continue
            m=CAP.match(s)
            if m and (m[1].lower() in ['表','table'])==(kind=='table'):cap=s;cap_index=j
            break
        result.append({'kind':kind,'body_index':i,'caption':cap,'caption_index':cap_index,'position_ok':cap is not None})
    return result

def audit_contract(path,contract,publication):
    d=Document(path);errors=[];checked=[];policy=normalize_publication_layout_policy(publication)
    def runcheck(p,r,fmt,loc,plain=False):
        v=effective_run(d,p,r);want={'ascii':fmt['latin'],'hAnsi':fmt['latin'],'eastAsia':fmt['east_asia'],'sz':str(round(fmt['size_pt']*2))}
        if plain:want.update(b=False,bCs=False)
        for k,w in want.items():
            if v.get(k)!=w:errors.append({'location':loc,'rule':k,'expected':w,'actual':v.get(k)})
        checked.append(loc)
    for i,p in enumerate(d.element.body.findall(qn('w:p'))):
        role=paragraph_role(d,p)
        if role not in ['body','caption','note'] or not visible(p) or not text(p).strip():continue
        fmt=contract['caption' if role=='caption' else 'body']
        props=effective_paragraph(d,p);sp=props.get('spacing',{})
        if sp.get('line')!=str(round(fmt['line_pt']*20)) or sp.get('lineRule')!='exact':
            errors.append({'location':f'p:{i}','rule':'body_caption_line_spacing'})
        ind=props.get('ind',{})
        expected_chars='200' if role=='body' else '0'
        if ind.get('firstLineChars')!=expected_chars:
            errors.append({'location':f'p:{i}','rule':'body_caption_firstLineChars','expected':expected_chars})
        for j,r in enumerate(p.xpath('.//w:r[w:t]')):
            if visible(r):runcheck(p,r,fmt,f'p:{i}:r:{j}',role=='caption')
    cellset=set();indset=set();alignset=set()
    for ti,t in enumerate(d.tables):
        if table_kind(t,policy)!='table':continue
        records=[r for r in _table_records(t) if visible(r['tc'])];sem=_column_semantics(records,policy)
        first_row=min(r['row'] for r in records);last_row=max(r['row'] for r in records)
        header_last=first_row+1 if '变量特定有效N' in text(t._tbl) else first_row
        for rr in records:
            tc=rr['tc'];loc=f't:{ti}:r:{rr["row"]}:c:{rr["cell"]}';cellset.add(loc)
            for side in ['top','bottom','left','right']:
                edge=tc.find('w:tcPr/w:tcBorders/w:'+side,tc.nsmap)
                size=12 if side=='top' and rr['row']==first_row or side=='bottom' and rr['row']==last_row else 4 if side=='bottom' and rr['row']==header_last else None
                if edge is None or edge.get(qn('w:val'))!=('single' if size else 'nil') or size and edge.get(qn('w:sz'))!=str(size):
                    errors.append({'location':loc,'rule':'three_line_border_'+side,'expected_size':size})
            cls=classify_cell_text(text(tc),policy);expected=_empty_cell_alignment(rr,sem) if cls=='empty' else cls
            for pi,p in enumerate(tc.findall(qn('w:p'))):
                if not visible(p):continue
                ind=p.find('w:pPr/w:ind',p.nsmap)
                # Explicit character zeros are required to suppress inherited Normal indents.
                for k in ['firstLineChars']:
                    if ind is None or ind.get(qn('w:'+k))!='0':errors.append({'location':loc,'rule':k,'expected':'0'})
                inherited=effective_paragraph(d,p).get('ind',{})
                for k in ['hangingChars','leftChars','rightChars','firstLine','hanging','left','right','start','end','startChars','endChars']:
                    if inherited.get(k,'0')!='0':errors.append({'location':loc,'rule':k,'actual':inherited.get(k),'expected':'0'})
                if ind is not None:
                    for k,v in ind.attrib.items():
                        if v!='0':errors.append({'location':loc,'rule':k,'actual':v,'expected':'0'})
                indset.add(loc)
                actual=effective_paragraph(d,p).get('jc',{}).get('val','left')
                if actual!=expected:errors.append({'location':loc,'rule':'alignment','expected':expected,'actual':actual})
                alignset.add(loc)
                sp=p.find('w:pPr/w:spacing',p.nsmap)
                expected_line_rule='atLeast' if p.xpath('.//m:oMath') and contract['table'].get('math_line_rule')=='atLeast' else 'exact'
                if sp is None or sp.get(qn('w:line'))!=str(round(contract['table']['line_pt']*20)) or sp.get(qn('w:lineRule'))!=expected_line_rule:errors.append({'location':loc,'rule':'line_spacing'})
                for j,r in enumerate(p.xpath('.//w:r[w:t]')):
                    if visible(r):runcheck(p,r,contract['table'],loc+f':p:{pi}:r:{j}',True)
    objects=object_registry(d,publication);caps=[o['caption'] for o in objects]
    if len(objects)!=contract['expected_objects']:errors.append({'rule':'object_count','actual':len(objects),'expected':contract['expected_objects']})
    for o in objects:
        if not o['position_ok']:errors.append({'rule':'caption_position','object':o})
    if len(set(caps))!=len(caps):errors.append({'rule':'duplicate_or_missing_caption'})
    cov=cellset==indset==alignset
    return {'all_pass':not errors and cov,'coverage_ok':cov,'table_cells':len(cellset),'runs_checked':len(checked),'objects':objects,'errors':errors,'scope':'accepted-view typography, effective run fonts, explicit unit/character zero indents, literal alignment, object-caption pairing'}
