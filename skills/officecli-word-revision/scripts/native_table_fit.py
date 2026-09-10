"""Native Word AutoFit plus tracked geometry transfer; never statistics."""
from pathlib import Path
import copy
import hashlib
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from manuscript_format import table_kind, visible, text, Recorder
from journal_layout import normalize_publication_layout_policy, canonicalize_property_order

def signature(table):
    return tuple(tuple(text(c).strip() for c in r.findall(qn("w:tc")))
                 for r in table._tbl.findall(qn("w:tr")) if visible(r))

def native_fit(run, path, policy):
    """run must be the pipeline's locked Runner.run; no foreground activation."""
    path=Path(path)
    p=normalize_publication_layout_policy(policy)
    mode=p["table"]["fit_mode"]
    d=Document(path)
    selected=[i+1 for i,t in enumerate(d.tables) if table_kind(t,p)=="table"]
    if mode=="preserve" or not selected:
        return {"mode":mode,"skipped":True,"tables":[]}
    if d.element.xpath(".//w:ins|.//w:del"):
        raise ValueError("Native table sizing requires an accepted-content candidate, not a redline")
    before=[signature(t) for t in d.tables]
    expr={"content":"auto fit content","window":"auto fit window","fixed":"auto fit fixed"}[mode]
    script=f'''tell application "Microsoft Word"
with timeout of 90 seconds
set priorAlerts to display alerts
try
set display alerts to alerts none
if exists document {json.dumps(path.name)} then error "Owned candidate is already open; do not close another writer's document"
open POSIX file {json.dumps(str(path))}
set docRef to document {json.dumps(path.name)}
if (count of tables of docRef) is not {len(d.tables)} then error "Word table index contract mismatch"
set track revisions of docRef to false
repeat with tableIndex in {{{",".join(map(str,selected))}}}
set tableRef to table (tableIndex as integer) of docRef
auto fit behavior tableRef behavior {expr}
end repeat
save docRef
close docRef saving yes
set display alerts to priorAlerts
return "WORDREV_NATIVE_AUTOFIT_COMPLETE:{mode}:{len(selected)}"
on error errMsg number errNum
set display alerts to priorAlerts
error errMsg number errNum
end try
end timeout
end tell'''
    result=run(["osascript","-e",script],timeout=110)
    marker=f"WORDREV_NATIVE_AUTOFIT_COMPLETE:{mode}:{len(selected)}"
    if marker not in result.stdout:raise RuntimeError("Word native AutoFit completion missing")
    after=Document(path)
    if before != [signature(t) for t in after.tables]:raise RuntimeError("Native AutoFit changed table text")
    report={"mode":mode,"native_completion":marker,"tables":selected,
            "skipped_diagram_tables":[i+1 for i,t in enumerate(d.tables) if table_kind(t,p)!="table"],
            "text_unchanged":True,"sha256":hashlib.sha256(path.read_bytes()).hexdigest()}
    return report

def transfer_geometry(prepared, tracked, policy, author, date):
    """Transfer Word-resolved widths into live redline rows as format revisions.

    Prior cumulative snapshots stay intact, so reject-all restores the mother.
    The native operation is performed once on the clean layout, never on deleted
    tables in a composite redline.
    """
    p=normalize_publication_layout_policy(policy)
    if p["table"]["fit_mode"]=="preserve":return []
    clean=Document(prepared);red=Document(tracked);rec=Recorder(red,author,date,True)
    matches={}
    for t in red.tables:
        if table_kind(t,p)=="table":matches.setdefault(signature(t),[]).append(t)
    def replace_children(dst,src,names):
        for name in names:
            tag=qn("w:"+name)
            for old in list(dst.findall(tag)):dst.remove(old)
            for new in src.findall(tag):dst.append(copy.deepcopy(new))
    for ci,ct in enumerate(clean.tables):
        if table_kind(ct,p)!="table":continue
        options=matches.get(signature(ct),[])
        if len(options)!=1:raise ValueError(f"Ambiguous/missing native-fit table match: {ci}")
        rt=options[0]
        rec.edit(rt._tbl,"w:tblPr",lambda pr:replace_children(pr,ct._tbl.tblPr,["tblW","tblLayout","tblInd","jc"]),f"native-fit:{ci}:table")
        oldgrid=rt._tbl.find(qn("w:tblGrid"));newgrid=ct._tbl.find(qn("w:tblGrid"))
        if newgrid is not None:
            if oldgrid is None:raise ValueError("Missing redline table grid")
            prior=copy.deepcopy(oldgrid)
            if [x.get(qn("w:w")) for x in oldgrid.findall(qn("w:gridCol"))] != [x.get(qn("w:w")) for x in newgrid.findall(qn("w:gridCol"))]:
                existing=oldgrid.find(qn("w:tblGridChange"))
                for x in list(oldgrid.findall(qn("w:gridCol"))):oldgrid.remove(x)
                for i,x in enumerate(newgrid.findall(qn("w:gridCol"))):oldgrid.insert(i,copy.deepcopy(x))
                if existing is None:
                    change=OxmlElement("w:tblGridChange")
                    change.set(qn("w:id"),str(rec.next));rec.next+=1
                    # CT_TblGridChange is CT_Markup: id only, unlike tblPrChange.
                    # Author/date are carried by adjacent property revisions and the run manifest.
                    for x in list(prior.findall(qn("w:tblGridChange"))):prior.remove(x)
                    change.append(prior);oldgrid.append(change)
        rr=[r for r in rt._tbl.findall(qn("w:tr")) if visible(r)]
        for ri,(a,b) in enumerate(zip(ct._tbl.findall(qn("w:tr")),rr)):
            ac=a.findall(qn("w:tc"));bc=b.findall(qn("w:tc"))
            if len(ac)!=len(bc):raise ValueError("Native-fit cell structure mismatch")
            for j,(ca,cb) in enumerate(zip(ac,bc)):
                rec.edit(cb,"w:tcPr",lambda pr,ca=ca:replace_children(pr,ca.find(qn("w:tcPr")),["tcW"]),f"native-fit:{ci}:{ri}:{j}")
    canonicalize_property_order(red);red.save(tracked)
    return rec.rows
