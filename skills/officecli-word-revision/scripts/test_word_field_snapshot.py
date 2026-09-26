import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from make_thesis_fixture import bookmark, element, field, make_fixture
from word_field_snapshot import compare, snapshot


class WordFieldSnapshotTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.path = self.root / "a.docx"
        self.contract = make_fixture(self.path)
        d = Document(self.path)
        bookmark(d.paragraphs[-1], "_Toc123", 100)
        # A duplicate visible title at another position is a deliberate trap.
        p = d.add_paragraph(d.paragraphs[-1].text)
        bookmark(p, "_Toc456", 101)
        p = d.add_paragraph()
        field(p, ' PAGEREF _Toc123 \\h ', "9")
        link = element("w:hyperlink", anchor="_Toc123")
        run, text = element("w:r"), element("w:t")
        text.text = "Appendix link"; run.append(text); link.append(run); p._p.append(link)
        d.save(self.path)
        self.first = snapshot(self.path, self.contract, 9)

    def second(self, mutate, pages=9):
        d = Document(self.path)
        mutate(d)
        p = self.root / "b.docx"
        d.save(p)
        return snapshot(p, self.contract, pages)

    def test_regenerated_names_require_same_resolved_location_and_text(self):
        def change(d):
            for n in d.element.iter():
                for key in (qn("w:name"), qn("w:anchor")):
                    if n.get(key) == "_Toc123":
                        n.set(key, "_Toc789")
                if n.tag == qn("w:instrText") and "_Toc123" in (n.text or ""):
                    n.text = n.text.replace("_Toc123", "_Toc789")
        result = compare(self.first, self.second(change))
        self.assertTrue(result["semantic_equal"], result)
        self.assertTrue(result["auto_toc_name_only_difference"])
        self.assertFalse(result["raw_equal"])

    def test_same_title_at_different_location_does_not_pass(self):
        def change(d):
            for n in d.element.iter(qn("w:instrText")):
                if "_Toc123" in (n.text or ""):
                    n.text = n.text.replace("_Toc123", "_Toc456")
        self.assertFalse(compare(self.first, self.second(change))["semantic_equal"])

    def test_materialized_run_properties_do_not_move_display_target(self):
        from docx.shared import Pt
        def change(d):
            for p in d.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(12)
        result = compare(self.first, self.second(change))
        self.assertTrue(result["semantic_equal"], result)

    def test_internal_link_changed_but_pageref_unchanged_fails(self):
        def change(d):
            next(d.element.iter(qn("w:hyperlink"))).set(qn("w:anchor"), "_Toc456")
        self.assertFalse(compare(self.first, self.second(change))["semantic_equal"])

    def test_missing_auto_target_is_not_discarded(self):
        def change(d):
            for n in d.element.iter(qn("w:instrText")):
                if "_Toc123" in (n.text or ""):
                    n.text = n.text.replace("_Toc123", "_Toc999")
        second = self.second(change)
        self.assertEqual(second["errors"][-1]["code"], "AUTO_TOC_TARGET_UNRESOLVED")
        self.assertFalse(compare(self.first, second)["semantic_equal"])

    def test_unchanged_files_with_different_page_counts_fail(self):
        self.assertFalse(compare(self.first, self.second(lambda d: None, pages=10))["semantic_equal"])

    def test_changed_cached_page_result_fails(self):
        def change(d):
            n = next(n for n in d.element.iter(qn("w:instrText")) if "_Toc123" in (n.text or ""))
            n.getparent().find(qn("w:t")).text = "10"
        self.assertFalse(compare(self.first, self.second(change))["semantic_equal"])


if __name__ == "__main__":
    unittest.main()
