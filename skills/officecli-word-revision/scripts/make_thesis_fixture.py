"""Generate an entirely synthetic thesis/contract pair for QA demonstrations.

Zotero-shaped fixture fields test parsing only; they are not native Zotero
insertion or refresh evidence. No real author, study, data or attachment is used.
"""
from __future__ import annotations

import argparse
import io
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw


def element(tag, **attrs):
    result = OxmlElement(tag)
    for key, value in attrs.items():
        result.set(qn("w:" + key), str(value))
    return result


def field(paragraph, code, result, *, simple=False):
    if simple:
        node = element("w:fldSimple", instr=code)
        run = element("w:r")
        text = element("w:t")
        text.text = result
        run.append(text)
        node.append(run)
        paragraph._p.append(node)
        return node
    run = paragraph.add_run()._r
    run.append(element("w:fldChar", fldCharType="begin"))
    instruction = element("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = code
    run.append(instruction)
    run.append(element("w:fldChar", fldCharType="separate"))
    text = element("w:t")
    text.text = result
    run.append(text)
    run.append(element("w:fldChar", fldCharType="end"))
    return run


def bookmark(paragraph, name, ident):
    start, end = element("w:bookmarkStart", id=ident, name=name), element("w:bookmarkEnd", id=ident)
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)
    return start, end


def bibliography_field(document, entries):
    """Synthetic, paragraph-spanning field; never a native Zotero claim."""
    paragraphs = [document.add_paragraph() for _ in entries]
    code = ' ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY '
    first = field(paragraphs[0], code, entries[0]["text"])
    end = first.find('w:fldChar[@w:fldCharType="end"]', first.nsmap)
    first.remove(end)
    for paragraph, entry in zip(paragraphs[1:], entries[1:]):
        paragraph.add_run(entry["text"])
    paragraphs[-1].add_run()._r.append(end)
    for paragraph in paragraphs:
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "0")


def make_fixture(path):
    d = Document()
    d.core_properties.author = "Synthetic QA fixture"
    d.core_properties.last_modified_by = "Synthetic QA fixture"
    d.core_properties.title = "Synthetic full thesis format example"
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)
    d.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
    d.styles["Normal"].paragraph_format.line_spacing = Pt(20)
    d.styles["Normal"].paragraph_format.first_line_indent = Pt(24)
    d.styles["Normal"].element.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")
    caption_style = d.styles.add_style("Synthetic Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption_style.base_style = d.styles["Normal"]
    caption_style.font.size = Pt(10.5)
    caption_style.paragraph_format.first_line_indent = Pt(0)
    caption_style.element.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "0")
    for name, size, east_asia in (("TOC 1", 14, "SimHei"), ("TOC 2", 12, "SimSun"),
                                  ("TOC 3", 12, "SimSun"), ("Footnote Text", 9, "SimSun")):
        style = d.styles[name] if name in d.styles else d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = d.styles["Normal"]
        style.font.name, style.font.size = "Times New Roman", Pt(size)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    # This minimal fixture has no recto-start requirement. Alternating stories
    # plus an odd page-number restart can produce a parity blank in Word PDF.
    d.settings.odd_and_even_pages_header_footer = False
    sec = d.sections[0]
    sec._sectPr.append(element("w:pgNumType", fmt="upperRoman", start="1"))
    sec.header.paragraphs[0].text = "Synthetic thesis"
    field(sec.footer.paragraphs[0], " PAGE ", "I", simple=True)
    field(sec.even_page_footer.paragraphs[0], " PAGE ", "II", simple=True)
    for heading, body in (("中文摘要", "这是仅用于测试的合成摘要。"),
                          ("Abstract", "This is a synthetic abstract for testing.")):
        p = d.add_paragraph(heading)
        p.paragraph_format.page_break_before = heading != "中文摘要"
        d.add_paragraph(body)
    for title, code, result in (("目录", ' TOC \\o "1-3" \\h ', "第1章 合成示例\t1"),
                                ("图目录", ' TOC \\c "图" \\h ', "图 1-1 合成图\t1"),
                                ("表目录", ' TOC \\c "表" \\h ', "表 1-1 合成表\t2")):
        p = d.add_paragraph(title)
        p.paragraph_format.page_break_before = True
        field(d.add_paragraph(), code, result)
    sec = d.add_section(WD_SECTION_START.NEW_PAGE)
    pg = sec._sectPr.find(qn("w:pgNumType"))
    pg.set(qn("w:fmt"), "decimal")
    pg.set(qn("w:start"), "1")
    d.add_paragraph("第1章 合成示例", "Heading 1")
    p = d.add_paragraph("合成论述见")
    field(p, " REF Figure11 \\h ", "图 1-1")
    p.add_run("和")
    field(p, " REF Table11 \\h ", "表 1-1", simple=True)
    p.add_run("；解释采用页底脚注。")
    reference = p.add_run()
    reference.font.superscript = True
    reference._r.append(element("w:footnoteReference", id="1"))
    item = {"citationID": "synthetic-only", "citationItems": [{"id": "fixture-a", "uris": ["https://example.org/synthetic/item-a"],
            "itemData": {"id": "fixture-a", "type": "book", "title": "Synthetic Methods", "author": [{"family": "Adams", "given": "A"}], "issued": {"date-parts": [[2020]]}}}], "properties": {"noteIndex": 0}}
    field(d.add_paragraph(), " ADDIN ZOTERO_ITEM CSL_CITATION " + json.dumps(item), "(Adams, 2020)")
    image = Image.new("RGB", (720, 300), "white")
    ImageDraw.Draw(image).rectangle((30, 30, 690, 270), outline="black", width=3)
    label = Image.new("RGBA", (240, 30))
    ImageDraw.Draw(label).text((0, 0), "SYNTHETIC FIGURE 1-1", fill="black")
    label = label.crop(label.getbbox())
    label = label.resize((label.width * 3, label.height * 3), Image.Resampling.NEAREST)
    image.paste(label, (80, 130), label)
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(io.BytesIO(stream.getvalue()), width=Cm(10))
    bookmark(p, "ObjectFigure11", 10)
    p = d.add_paragraph("图 1-", "Synthetic Caption")
    field(p, " SEQ 图 \\r 1 ", "1")
    p.add_run(" 合成图")
    p.paragraph_format.keep_with_next = False
    bookmark(p, "Figure11", 11)
    p = d.add_paragraph("表 1-", "Synthetic Caption")
    p.paragraph_format.keep_with_next = True
    field(p, " SEQ 表 \\r 1 ", "1")
    p.add_run(" 合成表")
    bookmark(p, "Table11", 12)
    table = d.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, (("指标", "Value"), ("合成 A", "1.25"), ("合成 B", "−0.50"))):
        for cell, value in zip(row.cells, values):
            p = cell.paragraphs[0]
            run = p.add_run(value)
            run.bold, run.font.size = False, Pt(10.5)
            props = p._p.get_or_add_pPr()
            props.append(element("w:jc", val="left" if "合成" in value or value == "指标" else "right"))
            props.append(element("w:ind", firstLine="0", hanging="0", firstLineChars="0", hangingChars="0"))
        row._tr.get_or_add_trPr().append(element("w:trHeight", val="280", hRule="atLeast"))
    table.rows[0]._tr.get_or_add_trPr().append(element("w:tblHeader"))
    bookmark(table.cell(0, 0).paragraphs[0], "ObjectTable11", 13)
    p = d.add_paragraph("本页合成说明用于验证脚注每页重新编号。")
    p.paragraph_format.page_break_before = True
    reference = p.add_run()
    reference.font.superscript = True
    reference._r.append(element("w:footnoteReference", id="2"))
    p = d.add_paragraph("参考文献")
    p.paragraph_format.page_break_before = True
    entries = [
        {"id": "en-adams", "language": "en", "author_display": "ADAMS A", "author_sort": "adams a", "sort_key_source": "synthetic fixture author definition", "year": "2020", "title": "Synthetic Methods", "title_sort": "synthetic methods", "text": "ADAMS A, 2020. Synthetic Methods[M]. Example Press."},
        {"id": "zh-chen", "language": "zh", "author_display": "陈甲", "author_sort": "chen jia", "sort_key_source": "synthetic fixture explicitly assigns chen jia", "year": "2021", "title": "合成研究", "title_sort": "he cheng yan jiu", "text": "陈甲, 2021. 合成研究[M]. 示例出版社."},
    ]
    bibliography_field(d, entries)
    sec = d.add_section(WD_SECTION_START.NEW_PAGE)
    sec._sectPr.find(qn("w:pgNumType")).attrib.pop(qn("w:start"), None)
    d.add_paragraph("附录 A 合成说明")
    d.add_paragraph("合成附录与版本化电子材料的引用示例。")
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(4.5), Cm(4)
        sec.left_margin, sec.right_margin = Cm(3.5), Cm(3)
        props = element("w:footnotePr")
        props.append(element("w:pos", val="pageBottom"))
        props.append(element("w:numRestart", val="eachPage"))
        # sectPr places note properties after header/footer references and
        # before page geometry. Do not append them after docGrid.
        index = next((i for i, node in enumerate(sec._sectPr)
                      if node.tag not in {qn("w:headerReference"), qn("w:footerReference")}),
                     len(sec._sectPr))
        sec._sectPr.insert(index, props)
    notes = element("w:footnotes")
    for ident, kind, payload in (("-1", "separator", "separator"), ("0", "continuationSeparator", "continuationSeparator")):
        note = element("w:footnote", id=ident, type=kind)
        p, r = element("w:p"), element("w:r")
        r.append(element("w:" + payload)); p.append(r); note.append(p); notes.append(note)
    for ident, content in (("1", "本注释完全为合成测试内容。"),
                           ("2", "第二个测试脚注与前注同节但不同页，应重新显示为1。")):
        note = element("w:footnote", id=ident)
        p, r, t = element("w:p"), element("w:r"), element("w:t")
        pp = element("w:pPr"); pp.append(element("w:pStyle", val="FootnoteText")); p.append(pp)
        r.append(element("w:footnoteRef")); t.text = content; r.append(t); p.append(r); note.append(p); notes.append(note)
    from lxml import etree
    part = Part(PackURI("/word/footnotes.xml"), "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", etree.tostring(notes), d.part.package)
    d.part.relate_to(part, RT.FOOTNOTES)
    props = element("w:footnotePr")
    props.append(element("w:pos", val="pageBottom")); props.append(element("w:numRestart", val="eachPage"))
    d.settings.element.insert_element_before(props, "w:endnotePr", "w:compat")
    d.save(path)
    blocks = [{"id": ident, "kind": kind, "anchor": {"text": title}, "start": "new_page"}
              for ident, kind, title in (("zh", "abstract_zh", "中文摘要"), ("en", "abstract_en", "Abstract"),
              ("toc", "contents", "目录"), ("lof", "figures", "图目录"), ("lot", "tables", "表目录"),
              ("ch1", "chapter", "第1章 合成示例"), ("bib", "references", "参考文献"), ("app", "appendix", "附录 A 合成说明"))]
    return {"schema_version": "1.0", "blocks": blocks,
            "sections": [{"section": 1, "page_number": {"format": "roman", "start": 1}},
                         {"section": 2, "page_number": {"format": "decimal", "start": 1}},
                         {"section": 3, "page_number": {"format": "decimal", "start": "continue"}, "stories": {"header_default": "linked"}}],
            "fields": {"unlocked": True, "minimum_zotero_items": 1, "zotero_bibliographies": 1,
                       "toc": {"contents": 1, "figures": 1, "tables": 1, "figure_label": "图", "table_label": "表"}},
            "footnotes": {"position": "pageBottom", "restart": "eachPage", "minimum_count": 2, "superscript": True},
            "captions": [{"id": kind + "11", "kind": kind, "label": label, "chapter": 1, "number": 1,
                          "caption": {"bookmark": name}, "object": {"bookmark": "Object" + name}}
                         for kind, label, name in (("figure", "图", "Figure11"), ("table", "表", "Table11"))],
            "caption_sequence_complete": True,
            "bibliography": {"heading": {"text": "参考文献"}, "group_order": ["en", "zh"],
                             "entries": entries, "entries_separate_paragraphs": True},
            "objects": {"inline_images": True, "tables": [{"anchor": {"bookmark": "ObjectTable11"}, "plain_header": True, "han_left_nonhan_right": True, "repeat_header": True}]}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--case", choices=("valid", "missing-reference", "appendix-restart",
                                           "bibliography-order", "clipped-image"), default="valid")
    args = parser.parse_args()
    root = Path(args.out_dir).resolve()
    root.mkdir(parents=True, exist_ok=False)
    contract = make_fixture(root / "synthetic-thesis.docx")
    expected_error = None
    if args.case != "valid":
        document = Document(root / "synthetic-thesis.docx")
        if args.case == "missing-reference":
            node = next(n for n in document.element.iter(qn("w:instrText")) if "REF Figure11" in (n.text or ""))
            node.text = " REF MissingTarget \\h "
            expected_error = "REFERENCE_TARGET_MISSING_OR_EMPTY"
        elif args.case == "appendix-restart":
            document.sections[-1]._sectPr.find(qn("w:pgNumType")).set(qn("w:start"), "1")
            expected_error = "PAGE_NUMBER_RESTART"
        elif args.case == "bibliography-order":
            node = next(n for n in document.element.iter(qn("w:instrText")) if "ZOTERO_BIBL" in (n.text or ""))
            node.getparent().find(qn("w:t")).text = "\n".join(e["text"] for e in reversed(contract["bibliography"]["entries"]))
            expected_error = "BIBLIOGRAPHY_VISIBLE_ORDER_OR_CONTENT"
        elif args.case == "clipped-image":
            paragraph = next(p for p in document.paragraphs if p._p.xpath(".//w:drawing"))
            paragraph.paragraph_format.line_spacing = Pt(20)
            expected_error = "IMAGE_EXACT_LINE_HEIGHT"
        document.save(root / "synthetic-thesis.docx")
    (root / "contract.json").write_text(json.dumps(contract, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    from thesis_integrity import audit
    result = audit(root / "synthetic-thesis.docx", contract)
    (root / "structural-qa.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    correct = result["all_pass"] if expected_error is None else expected_error in {e["code"] for e in result["errors"]}
    expectation = {"case": args.case, "expected_error": expected_error, "audit_status": result["status"],
                   "fixture_generation": "PASS" if correct else "FAIL", "native_word": "NOT_RUN", "native_zotero": "NOT_RUN"}
    (root / "expectation.json").write_text(json.dumps(expectation, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({**expectation, "output": str(root)}))
    # A negative fixture is successfully generated when its intended failure is
    # detected. The separate thesis_integrity CLI still exits 2 for that file.
    return 0 if correct else 2


if __name__ == "__main__":
    raise SystemExit(main())
