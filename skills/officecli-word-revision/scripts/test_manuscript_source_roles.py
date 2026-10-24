"""Behavioral controls for source-role checking, application and cache binding."""
import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

from manuscript_format import apply_contract, apply_source_roles, audit_contract, get, set_fonts, style_id, text
from manuscript_source_roles import SourceRoleError, load_bound_source
from test_manuscript_format import CONTRACT, PUB, fixture
from word_revision_pipeline import Pipeline, stable_hash


def pin(path):
    raw = path.read_bytes()
    return {"path": str(path), "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}


def digest(value):
    return hashlib.sha256(value.encode()).hexdigest()


def role_fixture():
    d = fixture()
    abstract = d.styles.add_style("Abstract", WD_STYLE_TYPE.PARAGRAPH)
    abstract.base_style = d.styles["Normal"]
    p = d.add_paragraph("摘要：甲", style=abstract)
    p.add_run(" italic")
    p.add_run("\t乙\n丙")
    d.add_paragraph("关键词：甲；乙", style=abstract)
    d.add_paragraph("参考文献")
    d.add_paragraph("Author A, 2020. First title.")
    d.add_paragraph("Author B, 2021. Second title.")
    d.add_paragraph("")
    apply_contract(d, CONTRACT, PUB)
    defaults = d.styles.element.find("w:docDefaults/w:pPrDefault/w:pPr", d.styles.element.nsmap)
    spacing = defaults.find(qn("w:spacing"))
    if spacing is not None:
        defaults.remove(spacing)
    for index in [2, 3, 5, 6]:
        p = d.paragraphs[index]
        pr = p._p.pPr
        pr.remove(pr.find(qn("w:spacing")))
        get(pr, "w:jc").set(qn("w:val"), "left")
        if index in [2, 3]:
            spacing = get(pr, "w:spacing")
            spacing.set(qn("w:before"), "100")
            spacing.set(qn("w:after"), "300")
        else:
            ind = get(pr, "w:ind")
            ind.attrib.clear()
            for key, value in {"firstLineChars": "0", "firstLine": "0",
                               "left": "720", "hanging": "720"}.items():
                ind.set(qn("w:" + key), value)
        fmt = dict(CONTRACT["body"], size_pt=10 if index in [2, 3] else 10.5)
        set_fonts(get(pr, "w:rPr"), fmt)
        for r in p.runs:
            set_fonts(get(r._r, "w:rPr"), fmt)
            if index in [5, 6]:
                get(r._r.rPr, "w:szCs").set(qn("w:val"), "24")
    get(d.paragraphs[2].runs[1]._r.rPr, "w:i").set(qn("w:val"), "1")
    return d


class SourceRoleTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="wordrev-source-roles-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "source.docx"
        role_fixture().save(self.source)
        d = Document(self.source)
        self.sidecar = {
            "schema_version": "1.0", "source": pin(self.source),
            "paragraph_count": len(d.paragraphs),
            "entries": [{"index": i, "role": role, "text_sha256": digest(text(d.paragraphs[i]._p)),
                         "style_id": style_id(d.paragraphs[i]._p)}
                        for i, role in [(2, "abstract"), (3, "keywords"),
                                        (5, "bibliography_entry"), (6, "bibliography_entry")]],
            "bibliography": {"heading_index": 4, "end_index": 7},
        }
        self.sidecar_path = self.root / "roles.json"
        self.contract = copy.deepcopy(CONTRACT)
        self.write_sidecar()
        self.output = self.root / "output.docx"

    def write_sidecar(self):
        self.sidecar_path.write_text(json.dumps(self.sidecar, ensure_ascii=False))
        self.contract["paragraph_roles"] = pin(self.sidecar_path)

    def audit(self, d):
        d.save(self.output)
        return audit_contract(self.output, self.contract, PUB)

    def good(self):
        d = Document(self.source)
        result = self.audit(d)
        self.assertTrue(result["all_pass"], result["errors"])
        self.assertEqual(result["source_roles"]["paragraphs_checked"], 4)
        self.assertEqual(result["table_cells"], 9)
        self.assertEqual(len(result["objects"]), 1)
        return d

    def test_special_roles_pass_without_exemptions_or_lost_table_coverage(self):
        d = self.good()
        legacy = audit_contract(self.output, CONTRACT, PUB)
        self.assertFalse(legacy["all_pass"])
        self.assertEqual(self.audit(d)["runs_checked"], legacy["runs_checked"])

    def test_equivalent_run_splitting_and_merging_pass(self):
        d = self.good()
        run = d.paragraphs[2].runs[0]._r
        value = run.find(qn("w:t")).text
        for part in [value[:2], value[2:]]:
            clone = copy.deepcopy(run)
            clone.find(qn("w:t")).text = part
            run.addprevious(clone)
        run.getparent().remove(run)
        self.assertTrue(self.audit(d)["all_pass"])
        first, second = d.paragraphs[2]._p.findall(qn("w:r"))[:2]
        first.find(qn("w:t")).text += second.find(qn("w:t")).text
        second.getparent().remove(second)
        self.assertTrue(self.audit(d)["all_pass"])

    def test_same_font_multiset_at_wrong_positions_fails(self):
        d = self.good()
        runs = d.paragraphs[2].runs
        get(runs[0]._r.rPr, "w:i").set(qn("w:val"), "1")
        get(runs[1]._r.rPr, "w:i").set(qn("w:val"), "0")
        self.assertFalse(self.audit(d)["all_pass"])

    def test_control_kinds_and_literal_whitespace_do_not_collapse(self):
        for kind in ["replace_tab", "remove_break", "literal_tab"]:
            with self.subTest(kind=kind):
                d = self.good()
                run = d.paragraphs[2].runs[2]._r
                if kind == "replace_tab":
                    run.find(qn("w:tab")).tag = qn("w:cr")
                elif kind == "remove_break":
                    run.remove(run.find(qn("w:br")))
                else:
                    tab = run.find(qn("w:tab"))
                    tab.tag = qn("w:t")
                    tab.text = "\t"
                self.assertFalse(self.audit(d)["all_pass"])

    def test_deleted_empty_paragraph_does_not_shift_roles(self):
        d = self.good()
        empty = OxmlElement("w:p")
        get(get(empty, "w:pPr"), "w:rPr").append(OxmlElement("w:del"))
        d.paragraphs[2]._p.addprevious(empty)
        result = self.audit(d)
        self.assertTrue(result["all_pass"], result["errors"])
        self.assertEqual(result["source_roles"]["checked"][0]["index"], 2)
        self.assertIn("p:3:", result["source_roles"]["checked"][0]["location"])

    def test_surviving_deleted_mark_merge_is_refused(self):
        d = self.good()
        get(get(d.paragraphs[0]._p, "w:pPr"), "w:rPr").append(OxmlElement("w:del"))
        with self.assertRaisesRegex(SourceRoleError, "MERGE_UNSUPPORTED"):
            self.audit(d)

    def test_font_size_cs_size_bold_and_paragraph_changes_fail(self):
        for change in ["size", "cs_size", "east_font", "latin_font", "bold",
                       "first_indent", "hanging", "line", "paragraph_spacing"]:
            with self.subTest(change=change):
                d = self.good()
                p = d.paragraphs[2 if change == "size" else 5]._p
                pr = p.find(qn("w:pPr"))
                rpr = p.find("w:r/w:rPr", p.nsmap)
                if change in ["size", "cs_size", "bold"]:
                    key, value = {"size": ("sz", "21"), "cs_size": ("szCs", "21"),
                                  "bold": ("b", "1")}[change]
                    get(rpr, "w:" + key).set(qn("w:val"), value)
                elif change in ["east_font", "latin_font"]:
                    get(rpr, "w:rFonts").set(qn("w:eastAsia" if change == "east_font" else "w:ascii"), "Arial")
                elif change in ["first_indent", "hanging"]:
                    get(pr, "w:ind").set(qn("w:firstLineChars" if change == "first_indent" else "w:hanging"), "200")
                else:
                    spacing = get(pr, "w:spacing")
                    spacing.set(qn("w:line" if change == "line" else "w:after"), "360")
                self.assertFalse(self.audit(d)["all_pass"])

    def test_source_apply_restores_multiple_roles_and_keeps_controls(self):
        d = self.good()
        get(d.paragraphs[2].runs[0]._r.rPr, "w:sz").set(qn("w:val"), "21")
        get(d.paragraphs[5]._p.pPr, "w:ind").set(qn("w:hanging"), "0")
        self.assertFalse(self.audit(d)["all_pass"])
        apply_contract(d, self.contract, PUB, tracked=True)
        self.assertTrue(self.audit(d)["all_pass"])
        self.assertEqual(d.paragraphs[5].runs[0]._r.rPr.find(qn("w:szCs")).get(qn("w:val")), "24")
        self.assertTrue(d.paragraphs[2]._p.xpath(".//w:tab"))
        self.assertTrue(d.element.xpath(".//w:rPrChange|.//w:pPrChange"))

    def test_minimal_source_apply_leaves_body_table_and_already_good_roles_untouched(self):
        d = self.good()
        before = d.element.xml
        self.assertEqual(apply_source_roles(d, self.contract), [])
        self.assertEqual(d.element.xml, before)
        body = d.paragraphs[0]._p.xml
        table = d.tables[0]._tbl.xml
        get(d.paragraphs[5]._p.pPr, "w:ind").set(qn("w:hanging"), "0")
        apply_source_roles(d, self.contract, tracked=True)
        self.assertEqual(d.paragraphs[0]._p.xml, body)
        self.assertEqual(d.tables[0]._tbl.xml, table)
        self.assertTrue(self.audit(d)["all_pass"])

    def test_changed_inherited_abstract_size_is_not_hidden_by_the_role(self):
        d = self.good()
        for run in d.paragraphs[2].runs:
            pr = run._r.rPr
            pr.remove(pr.find(qn("w:sz")))
        get(get(d.styles["Abstract"].element, "w:rPr"), "w:sz").set(qn("w:val"), "24")
        self.assertFalse(self.audit(d)["all_pass"])

    def test_apply_can_split_a_plain_merged_run_at_source_font_boundaries(self):
        d = self.good()
        first, second = d.paragraphs[2]._p.findall(qn("w:r"))[:2]
        first.find(qn("w:t")).text += second.find(qn("w:t")).text
        second.getparent().remove(second)
        self.assertFalse(self.audit(d)["all_pass"])
        apply_contract(d, self.contract, PUB)
        self.assertTrue(self.audit(d)["all_pass"])

    def test_source_or_sidecar_drift_fails_closed(self):
        for which in ["source", "sidecar"]:
            with self.subTest(which=which):
                path = self.source if which == "source" else self.sidecar_path
                before = path.read_bytes()
                path.write_bytes(before + b" ")
                with self.assertRaisesRegex(SourceRoleError, "PIN_MISMATCH"):
                    load_bound_source(self.contract["paragraph_roles"])
                path.write_bytes(before)

    def test_bad_selectors_and_role_spoofing_fail_closed(self):
        original = copy.deepcopy(self.sidecar)
        for case in ["unknown", "duplicate", "out_of_range", "empty", "missing_abstract",
                     "missing_bib", "swap_labels", "body_as_bib", "unknown_key", "bad_style"]:
            with self.subTest(case=case):
                self.sidecar = copy.deepcopy(original)
                e = self.sidecar["entries"]
                if case == "unknown": e[0]["role"] = "body"
                elif case == "duplicate": e.append(copy.deepcopy(e[0]))
                elif case == "out_of_range": e[0]["index"] = 1000
                elif case == "empty": e.clear()
                elif case == "missing_abstract": e.pop(0)
                elif case == "missing_bib": e.pop()
                elif case == "swap_labels": e[0]["role"], e[1]["role"] = e[1]["role"], e[0]["role"]
                elif case == "body_as_bib":
                    body = Document(self.source).paragraphs[0]
                    e[0].update(index=0, role="bibliography_entry",
                                text_sha256=digest(text(body._p)), style_id=style_id(body._p))
                elif case == "unknown_key": e[0]["skip_errors"] = True
                else: e[0]["style_id"] = "Normal"
                self.write_sidecar()
                with self.assertRaises(SourceRoleError):
                    load_bound_source(self.contract["paragraph_roles"])
        self.sidecar = original
        self.write_sidecar()

    def test_text_and_style_drift_are_real_role_errors(self):
        for kind in ["text", "style"]:
            with self.subTest(kind=kind):
                d = self.good()
                if kind == "text": d.paragraphs[2].runs[0].text += "!"
                else: d.paragraphs[2].style = d.styles["Normal"]
                self.assertFalse(self.audit(d)["all_pass"])

    def test_nonselected_body_table_and_caption_are_still_checked(self):
        for kind in ["body", "table", "caption"]:
            with self.subTest(kind=kind):
                d = self.good()
                if kind == "body":
                    get(d.paragraphs[0].runs[0]._r.rPr, "w:sz").set(qn("w:val"), "40")
                elif kind == "table":
                    get(d.tables[0].cell(1, 0).paragraphs[0]._p.pPr, "w:jc").set(qn("w:val"), "center")
                else:
                    p = d.paragraphs[1]._p
                    p.getparent().remove(p)
                    # Add an ordinary blank to retain the role coordinate count.
                    d.element.body.insert(1, OxmlElement("w:p"))
                self.assertFalse(self.audit(d)["all_pass"])

    def job(self):
        adapter = self.root / "adapter.py"
        adapter.write_text("def configure(job, run_dir):\n return {}\n")
        job = {"schema_version": "2.0", "job_id": "roles", "scenario": "manuscript_revision",
               "source": str(self.source), "content_source": str(self.source), "adapter": str(adapter),
               "run_dir": str(self.root / "run"), "cache_dir": str(self.root / "cache"),
               "notebooklm": "disabled", "statistics": {"allow_execution": False, "consume_manifest_only": True},
               "format_contract": self.contract}
        path = self.root / "job.json"
        path.write_text(json.dumps(job))
        return path

    def test_pipeline_refuses_bad_binding_before_creating_run_or_loading_adapter(self):
        path = self.job()
        self.sidecar_path.write_bytes(self.sidecar_path.read_bytes() + b" ")
        with self.assertRaisesRegex(SourceRoleError, "PIN_MISMATCH"):
            Pipeline(path)
        self.assertFalse((self.root / "run").exists())

    def test_same_bytes_at_different_content_source_path_are_not_authority(self):
        path = self.job()
        clone = self.root / "other-source.docx"
        clone.write_bytes(self.source.read_bytes())
        job = json.loads(path.read_text())
        job["content_source"] = str(clone)
        path.write_text(json.dumps(job))
        with self.assertRaisesRegex(SourceRoleError, "CONTENT_SOURCE_PATH_MISMATCH"):
            Pipeline(path)

    def test_changed_helper_invalidates_cached_phase_and_input_drift_cannot_use_cache(self):
        pipeline = Pipeline(self.job())
        import word_revision_pipeline as module
        script_dir = self.root / "private_scripts"
        ref_dir = self.root / "references"
        script_dir.mkdir()
        ref_dir.mkdir()
        helper = script_dir / "manuscript_source_roles.py"
        helper.write_text("# original helper")
        (ref_dir / "manuscript-source-roles.schema.json").write_text("{}")
        (ref_dir / "job.schema.json").write_text("{}")
        output = self.root / "phase-output"
        calls = []
        def action():
            calls.append(1)
            output.write_text("unchanged output")
            return [output]
        with patch.object(module, "SCRIPT_DIR", script_dir), patch.object(module, "SKILL_DIR", self.root):
            def run():
                fingerprint = stable_hash({"inputs": pipeline._input_hashes(),
                                           "implementation": pipeline._implementation_hashes()})
                return pipeline._phase("audit", fingerprint, action, resume=True)
            run()
            self.assertTrue(run()["cache_hit"])
            helper.write_text("# changed helper")
            run()
            self.assertEqual(len(calls), 2)
            self.source.write_bytes(self.source.read_bytes() + b" ")
            with self.assertRaisesRegex(SourceRoleError, "PIN_MISMATCH"):
                run()
            self.assertEqual(len(calls), 2)

    def refresh_source(self, d):
        d.save(self.source)
        self.sidecar["source"] = pin(self.source)
        self.sidecar["paragraph_count"] = len(d.paragraphs)
        self.write_sidecar()

    def test_every_source_run_boolean_and_complex_attribute_change_is_detected(self):
        changes = [
            ("strike", {"val": "1"}), ("dstrike", {"val": "1"}),
            ("caps", {"val": "1"}), ("smallCaps", {"val": "1"}),
            ("outline", {"val": "1"}), ("shadow", {"val": "1"}),
            ("emboss", {"val": "1"}), ("imprint", {"val": "1"}),
            ("vanish", {"val": "1"}), ("webHidden", {"val": "1"}),
            ("rtl", {"val": "1"}), ("cs", {"val": "1"}),
            ("highlight", {"val": "yellow"}),
            ("shd", {"val": "clear", "fill": "FFFF00"}),
            ("u", {"val": "single", "color": "FF0000"}),
            ("lang", {"eastAsia": "ja-JP"}),
        ]
        for tag, attrs in changes:
            with self.subTest(property=tag):
                d = self.good()
                pr = get(d.paragraphs[2].runs[0]._r, "w:rPr")
                prop = get(pr, "w:" + tag)
                for name, value in attrs.items():
                    prop.set(qn("w:" + name), value)
                result = self.audit(d)
                self.assertFalse(result["all_pass"])
                self.assertTrue(any(e["rule"] == "source_role_run_format" for e in result["errors"]))

    def test_shading_and_underline_keep_attributes_beyond_val(self):
        d = Document(self.source)
        pr = get(d.paragraphs[2].runs[0]._r, "w:rPr")
        get(pr, "w:shd").set(qn("w:val"), "clear")
        get(pr, "w:shd").set(qn("w:fill"), "FFFF00")
        get(pr, "w:u").set(qn("w:val"), "single")
        get(pr, "w:u").set(qn("w:color"), "FF0000")
        self.refresh_source(d)
        for prop, attr, value in [("shd", "fill", "00FF00"), ("u", "color", "0000FF")]:
            with self.subTest(property=prop):
                candidate = self.good()
                get(candidate.paragraphs[2].runs[0]._r.rPr, "w:" + prop).set(qn("w:" + attr), value)
                self.assertFalse(self.audit(candidate)["all_pass"])

    def test_unknown_run_extension_is_checked_and_direct_changes_can_be_restored(self):
        from lxml import etree
        d = self.good()
        extension = etree.Element("{urn:source-role-future-format}effect")
        extension.set("strength", "3")
        get(d.paragraphs[2].runs[0]._r, "w:rPr").append(extension)
        self.assertFalse(self.audit(d)["all_pass"])
        edits = apply_source_roles(d, self.contract)
        self.assertTrue(edits)
        self.assertTrue(self.audit(d)["all_pass"])

    def test_equivalent_false_spellings_and_style_toggles_keep_effective_format(self):
        d = self.good()
        get(d.paragraphs[2].runs[0]._r.rPr, "w:strike").set(qn("w:val"), "off")
        self.assertTrue(self.audit(d)["all_pass"])
        d = Document(self.source)
        get(get(d.styles["Abstract"].element, "w:rPr"), "w:strike").set(qn("w:val"), "1")
        char = d.styles.add_style("StrikeToggle", WD_STYLE_TYPE.CHARACTER)
        get(get(char.element, "w:rPr"), "w:strike").set(qn("w:val"), "true")
        get(d.paragraphs[2].runs[0]._r.rPr, "w:rStyle").set(qn("w:val"), char.style_id)
        self.refresh_source(d)
        candidate = self.good()
        pr = candidate.paragraphs[2].runs[0]._r.rPr
        pr.remove(pr.find(qn("w:rStyle")))
        get(pr, "w:strike").set(qn("w:val"), "false")
        self.assertTrue(self.audit(candidate)["all_pass"])

    def test_inherited_and_paragraph_mark_effects_are_detected(self):
        for scope in ["style", "mark"]:
            with self.subTest(scope=scope):
                d = self.good()
                pr = get(d.styles["Abstract"].element, "w:rPr") if scope == "style" else (
                    get(get(d.paragraphs[2]._p, "w:pPr"), "w:rPr"))
                get(pr, "w:caps").set(qn("w:val"), "1")
                self.assertFalse(self.audit(d)["all_pass"])

    def test_duplicate_or_invalid_onoff_run_properties_are_refused(self):
        for case in ["duplicate", "invalid"]:
            with self.subTest(case=case):
                d = self.good()
                pr = d.paragraphs[2].runs[0]._r.rPr
                get(pr, "w:strike").set(qn("w:val"), "1" if case == "duplicate" else "not-a-boolean")
                if case == "duplicate":
                    pr.append(OxmlElement("w:strike"))
                with self.assertRaisesRegex(SourceRoleError, "DUPLICATE_RUN_PROPERTY|INVALID_ONOFF"):
                    self.audit(d)

    def test_table_cell_labels_fail_with_an_explicit_unsupported_coordinate(self):
        for label in ["摘要：表格摘要", "关键词：表格词", "参考文献"]:
            with self.subTest(label=label):
                d = Document(self.source)
                d.tables[0].cell(0, 0).paragraphs[0].text = label
                self.refresh_source(d)
                with self.assertRaisesRegex(SourceRoleError, "LABEL_IN_TABLE_UNSUPPORTED"):
                    load_bound_source(self.contract["paragraph_roles"])
        # A new nested label in the candidate is refused as well.
        self.refresh_source(role_fixture())
        d = self.good()
        d.tables[0].cell(0, 0).paragraphs[0].text = "关键词：后加"
        with self.assertRaisesRegex(SourceRoleError, "LABEL_IN_TABLE_UNSUPPORTED"):
            self.audit(d)

    def test_deleted_table_label_does_not_spoof_accepted_role_coverage(self):
        d = Document(self.source)
        d.tables[0].cell(0, 0).paragraphs[0].text = "摘要：已删除"
        row = d.tables[0].rows[0]._tr
        get(row, "w:trPr").append(OxmlElement("w:del"))
        self.refresh_source(d)
        load_bound_source(self.contract["paragraph_roles"])

    def test_other_nested_label_coordinates_are_not_silently_skipped(self):
        d = Document(self.source)
        sdt = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "摘要：内容控件"
        r.append(t); p.append(r); content.append(p); sdt.append(content)
        d.element.body.insert(1, sdt)
        self.refresh_source(d)
        with self.assertRaisesRegex(SourceRoleError, "LABEL_IN_NESTED_UNSUPPORTED"):
            load_bound_source(self.contract["paragraph_roles"])

    def test_bibliography_end_cannot_hide_a_real_trailing_entry(self):
        self.sidecar["bibliography"]["end_index"] = 6
        self.sidecar["entries"] = [e for e in self.sidecar["entries"] if e["index"] != 6]
        self.write_sidecar()
        with self.assertRaisesRegex(SourceRoleError, "END_TRUNCATES_CONTENT"):
            load_bound_source(self.contract["paragraph_roles"])

    def test_bibliography_tail_allows_blanks_or_a_real_next_heading(self):
        load_bound_source(self.contract["paragraph_roles"])
        d = Document(self.source)
        d.add_paragraph("附录 A", style="Heading 1")
        d.add_paragraph("Following section text.")
        self.refresh_source(d)
        load_bound_source(self.contract["paragraph_roles"])
        # An ordinary non-empty tail, after any blanks, is not a boundary.
        d.paragraphs[8].style = d.styles["Normal"]
        d.paragraphs[8].text = "Author C, 2022. Third title."
        self.refresh_source(d)
        with self.assertRaisesRegex(SourceRoleError, "END_TRUNCATES_CONTENT"):
            load_bound_source(self.contract["paragraph_roles"])

    def test_schema_accepts_only_the_narrow_role_pointer(self):
        import jsonschema
        schema = json.loads((Path(__file__).resolve().parent.parent / "references/job.schema.json").read_text())
        job = json.loads(self.job().read_text())
        jsonschema.validate(job, schema)
        job["format_contract"]["paragraph_roles"]["skip_errors"] = [1]
        with self.assertRaises(jsonschema.ValidationError):
            jsonschema.validate(job, schema)
        jsonschema.validate(self.sidecar, json.loads(
            (Path(__file__).resolve().parent.parent / "references/manuscript-source-roles.schema.json").read_text()))


if __name__ == "__main__":
    unittest.main()
