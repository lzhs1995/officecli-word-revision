#!/usr/bin/env python3
"""Generic journal table/layout normalization and rendered-PDF QA.

The helper is opt-in through ``job.publication_layout``.  It deliberately
leaves diagram-carrier tables alone because several legacy manuscripts use
tables as drawing canvases rather than as data tables.
"""

from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import subprocess
from collections import Counter
from pathlib import Path
from statistics import median
from typing import Any, Callable

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


HAN_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
ARROW_RE = re.compile(r"[\u2190-\u21ff\u27f0-\u27ff\u2900-\u297f]")
CAPTION_RE = re.compile(r"^(?:图|表|Figure|Table)\s*[0-9]", re.IGNORECASE)
NOTE_RE = re.compile(r"^(?:注|Note)\s*[:：]", re.IGNORECASE)
HEADING_RE = re.compile(r"^(?:[一二三四五六七八九十]+[、.]|[（(][一二三四五六七八九十0-9]+[）)]|[0-9]+(?:\.[0-9]+){0,3}\s)")


DEFAULT_PUBLICATION_LAYOUT: dict[str, Any] = {
    "mode": "generic_journal",
    "figure": {
        "max_width_ratio": 0.90,
        "caption_block_cm": 0.75,
        "narrative_lines": 2,
        "narrative_line_cm": 0.48,
        "min_label_pt_advisory": 8.0,
    },
    "table": {
        "fit_mode": "content",
        # Dissertation table rule: Han text left; all non-Han content
        # (including Latin letters and numbers) right. No prose exception.
        "alignment_mode": "han-left-nonhan-right",
        "prose_min_chars": 25,
        "prose_min_spaces": 3,
        "diagram_arrow_threshold": 3,
        "diagram_table_majority": True,
        "small_font_max_pt": 10.0,
        "small_line_spacing_pt": 12.0,
        "large_line_spacing_pt": 14.0,
        "cell_margin_twips": 102,
        "repeat_header": True,
        "cant_split_rows": True,
    },
    "pagination": {
        "require_narrative_on_object_pages": True,
        "minimum_narrative_lines": 2,
        "minimum_narrative_characters": 40,
        "maximum_repair_cycles": 2,
        "object_scope": None,
    },
}

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
    "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange",
]
TRPR_ORDER = [
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
    "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
    "ins", "del", "trPrChange",
]
TBLPR_ORDER = [
    'tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize',
    'tblStyleColBandSize','tblW','jc','tblCellSpacing','tblInd','tblBorders',
    'shd','tblLayout','tblCellMar','tblLook','tblCaption','tblDescription','tblPrChange',
]


def _deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = copy.deepcopy(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = _deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def default_publication_layout_policy() -> dict[str, Any]:
    return copy.deepcopy(DEFAULT_PUBLICATION_LAYOUT)


def normalize_publication_layout_policy(value: dict[str, Any] | None) -> dict[str, Any]:
    if value is None:
        value = {}
    if not isinstance(value, dict):
        raise ValueError("publication_layout must be an object")
    policy = _deep_merge(DEFAULT_PUBLICATION_LAYOUT, value)
    if policy["mode"] not in {"generic_journal", "source_master", "explicit_venue"}:
        raise ValueError("publication_layout.mode must be generic_journal, source_master, or explicit_venue")
    figure = policy["figure"]
    table = policy["table"]
    if table["fit_mode"] not in {"content", "window", "fixed", "preserve"}:
        raise ValueError("unsupported publication_layout.table.fit_mode")
    pagination = policy["pagination"]
    if not 0 < float(figure["max_width_ratio"]) <= 1:
        raise ValueError("publication_layout.figure.max_width_ratio must be in (0, 1]")
    for key in ("caption_block_cm", "narrative_line_cm", "min_label_pt_advisory"):
        if float(figure[key]) <= 0:
            raise ValueError(f"publication_layout.figure.{key} must be positive")
    if int(figure["narrative_lines"]) < 1:
        raise ValueError("publication_layout.figure.narrative_lines must be positive")
    if table["alignment_mode"] not in {"han-left-short-nonhan-right", "han-left-nonhan-right"}:
        raise ValueError("unsupported publication_layout.table.alignment_mode")
    for key in ("prose_min_chars", "prose_min_spaces", "diagram_arrow_threshold", "cell_margin_twips"):
        if int(table[key]) < 0:
            raise ValueError(f"publication_layout.table.{key} must be non-negative")
    for key in ("small_font_max_pt", "small_line_spacing_pt", "large_line_spacing_pt"):
        if float(table[key]) <= 0:
            raise ValueError(f"publication_layout.table.{key} must be positive")
    for key in ("minimum_narrative_lines", "minimum_narrative_characters", "maximum_repair_cycles"):
        if int(pagination[key]) < 0:
            raise ValueError(f"publication_layout.pagination.{key} must be non-negative")
    scope = pagination.get("object_scope")
    if scope is not None:
        if not isinstance(scope, dict):
            raise ValueError("publication_layout.pagination.object_scope must be an object or null")
        if not str(scope.get("start_text", "")).strip() or not str(scope.get("end_text", "")).strip():
            raise ValueError("publication_layout.pagination.object_scope requires start_text and end_text")
    return policy


def _text(element) -> str:
    return "".join(node.text or "" for node in element.xpath(".//w:t"))


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def _arrow_count(value: str) -> int:
    return len(ARROW_RE.findall(value or ""))


def _is_english_prose(value: str, policy: dict[str, Any]) -> bool:
    stripped = value.strip()
    table = policy["table"]
    return (
        not HAN_RE.search(stripped)
        and len(stripped) >= int(table["prose_min_chars"])
        and stripped.count(" ") >= int(table["prose_min_spaces"])
    )


def classify_cell_text(value: str, policy: dict[str, Any]) -> str:
    """Return left/right/diagram/empty for one visible cell string."""
    stripped = value.strip()
    if not stripped:
        return "empty"
    if _arrow_count(stripped) >= int(policy["table"]["diagram_arrow_threshold"]):
        return "diagram"
    if HAN_RE.search(stripped) or (policy["table"]["alignment_mode"] != "han-left-nonhan-right" and _is_english_prose(stripped, policy)):
        return "left"
    return "right"


def _grid_span(tc) -> int:
    node = tc.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
    if node is None:
        return 1
    try:
        return max(1, int(node.get(qn("w:val"), "1")))
    except ValueError:
        return 1


def _table_records(table) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for row_index, tr in enumerate(table._tbl.findall(qn("w:tr"))):
        grid_column = 0
        for cell_index, tc in enumerate(tr.findall(qn("w:tc"))):
            span = _grid_span(tc)
            records.append({
                "row": row_index,
                "cell": cell_index,
                "grid_start": grid_column,
                "grid_end": grid_column + span,
                "span": span,
                "tc": tc,
                "text": _text(tc).strip(),
            })
            grid_column += span
    return records


def _table_is_diagram(records: list[dict[str, Any]], policy: dict[str, Any]) -> bool:
    nonempty = [row for row in records if row["text"]]
    if not nonempty:
        return False
    diagram = sum(classify_cell_text(row["text"], policy) == "diagram" for row in nonempty)
    return bool(policy["table"]["diagram_table_majority"]) and diagram * 2 >= len(nonempty)


def _doc_default_font_pt(document: Document) -> float:
    nodes = document.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:sz")
    if nodes:
        return float(nodes[0].get(qn("w:val"))) / 2.0
    return 10.5


def _style_font_half_points(document: Document, style_id: str) -> list[int]:
    if not style_id:
        return []
    seen: set[str] = set()
    current = style_id
    while current and current not in seen:
        seen.add(current)
        rows = document.styles.element.xpath(f"./w:style[@w:styleId='{current}']")
        if not rows:
            break
        style = rows[0]
        sizes = [int(node.get(qn("w:val"))) for node in style.xpath("./w:rPr/w:sz")]
        if sizes:
            return sizes
        based = style.find(qn("w:basedOn"))
        current = based.get(qn("w:val")) if based is not None else ""
    return []


def _table_font_pt(document: Document, records: list[dict[str, Any]]) -> float:
    direct: list[int] = []
    styled: list[int] = []
    for record in records:
        tc = record["tc"]
        direct.extend(
            int(node.get(qn("w:val")))
            for node in tc.xpath(".//w:r/w:rPr/w:sz")
            if node.get(qn("w:val"), "").isdigit()
        )
        for paragraph in tc.findall(qn("w:p")):
            ppr = paragraph.find(qn("w:pPr"))
            style = ppr.find(qn("w:pStyle")) if ppr is not None else None
            if style is not None:
                styled.extend(_style_font_half_points(document, style.get(qn("w:val"), "")))
    values = direct or styled
    return float(median(values)) / 2.0 if values else _doc_default_font_pt(document)


def _line_spacing_twips(font_pt: float, policy: dict[str, Any]) -> int:
    table = policy["table"]
    target = (
        float(table["small_line_spacing_pt"])
        if font_pt <= float(table["small_font_max_pt"])
        else float(table["large_line_spacing_pt"])
    )
    return int(round(target * 20))


def _column_semantics(records: list[dict[str, Any]], policy: dict[str, Any]) -> dict[int, str]:
    votes: dict[int, Counter] = {}
    for record in records:
        classification = classify_cell_text(record["text"], policy)
        if classification not in {"left", "right"}:
            continue
        for column in range(record["grid_start"], record["grid_end"]):
            votes.setdefault(column, Counter())[classification] += 1
    result: dict[int, str] = {}
    for column, counter in votes.items():
        left, right = counter["left"], counter["right"]
        if left == right:
            raise RuntimeError(f"grid column {column} has tied left/right semantics")
        result[column] = "left" if left > right else "right"
    return result


def _empty_cell_alignment(record: dict[str, Any], semantics: dict[int, str]) -> str:
    values = {
        semantics.get(column)
        for column in range(record["grid_start"], record["grid_end"])
    }
    values.discard(None)
    if len(values) != 1:
        raise RuntimeError(
            "empty cell has undefined or conflicting grid-column semantics: "
            f"row={record['row']} cell={record['cell']} columns="
            f"{record['grid_start']}:{record['grid_end']}"
        )
    return values.pop()


def _get_or_add(parent, tag: str):
    node = parent.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        parent.append(node)
    return node


def _set_cell_margin(tc, twips: int) -> None:
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.insert_element_before(
            margins, "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"
        )
    for side in ("left", "right"):
        node = _get_or_add(margins, f"w:{side}")
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")


def _set_paragraph_format(paragraph, alignment: str, line_twips: int) -> None:
    ppr = paragraph.get_or_add_pPr()
    jc = ppr.get_or_add_jc()
    jc.set(qn("w:val"), alignment)
    ind = ppr.get_or_add_ind()
    for name in ("left", "right", "start", "end", "firstLine", "hanging"):
        ind.set(qn(f"w:{name}"), "0")
    for name in ("leftChars", "rightChars", "startChars", "endChars", "firstLineChars", "hangingChars"):
        # Absence inherits a parent style's character-unit indent. It is not
        # equivalent to zero: Word can reinstate two-character indentation.
        ind.set(qn(f"w:{name}"), "0")
    spacing = ppr.get_or_add_spacing()
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), str(line_twips))
    spacing.set(qn("w:lineRule"), "exact")


def _set_row_flags(table, policy: dict[str, Any], *, diagram: bool) -> None:
    rows = table._tbl.findall(qn("w:tr"))
    if diagram:
        return
    for tr in rows:
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            tr.insert(0, trpr)
        if policy["table"]["cant_split_rows"]:
            cant_split = trpr.find(qn("w:cantSplit"))
            if cant_split is None:
                cant_split = OxmlElement("w:cantSplit")
                trpr.insert_element_before(
                    cant_split, "w:trHeight", "w:tblHeader", "w:tblCellSpacing",
                    "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange"
                )
    if len(rows) > 1 and policy["table"]["repeat_header"]:
        first_pr = rows[0].find(qn("w:trPr"))
        header = first_pr.find(qn("w:tblHeader"))
        if header is None:
            header = OxmlElement("w:tblHeader")
            first_pr.insert_element_before(
                header, "w:tblCellSpacing", "w:jc", "w:hidden",
                "w:ins", "w:del", "w:trPrChange"
            )


def _xml_sha256(element) -> str:
    return hashlib.sha256(element.xml.encode("utf-8")).hexdigest()


def _sort_children(element, order: list[str]) -> bool:
    rank = {name: index for index, name in enumerate(order)}
    children = list(element)
    decorated = []
    for position, child in enumerate(children):
        local = child.tag.rsplit("}", 1)[-1]
        decorated.append((rank.get(local, len(order)), position, child))
    sorted_children = [row[2] for row in sorted(decorated, key=lambda row: (row[0], row[1]))]
    if children == sorted_children:
        return False
    for child in children:
        element.remove(child)
    for child in sorted_children:
        element.append(child)
    return True


def canonicalize_property_order(document: DocumentClass) -> dict[str, int]:
    """Restore WordprocessingML property sequence after external editors."""
    changed = {"pPr": 0, "tcPr": 0, "trPr": 0}
    for name, xpath, order in (
        ("pPr", ".//w:pPr", PPR_ORDER),
        ("tcPr", ".//w:tcPr", TCPR_ORDER),
        ("trPr", ".//w:trPr", TRPR_ORDER),
        ("tblPr", ".//w:tblPr", TBLPR_ORDER),
    ):
        for element in document.element.xpath(xpath):
            changed[name] = changed.get(name,0) + int(_sort_children(element, order))
    # OfficeCLI 1.0.147 can place w14 ligatures after rPrChange while
    # serializing settings. Extension properties must precede that final item.
    changed['rPrChange_last']=0
    for pr in document.element.xpath('.//w:rPr'):
        change=pr.find(qn('w:rPrChange'))
        if change is not None and pr[-1] is not change:
            pr.remove(change);pr.append(change);changed['rPrChange_last']+=1
    return changed


def canonicalize_paragraph_ids(document: DocumentClass) -> dict[str, int]:
    """Keep Word 2010 paragraph/text IDs unique and below 0x80000000."""
    result = {"paraId": 0, "textId": 0}
    for local_name in ("paraId", "textId"):
        attribute = qn(f"w14:{local_name}")
        used: set[int] = set()
        pending = []
        for paragraph in document.element.xpath(".//w:p"):
            value = paragraph.get(attribute)
            try:
                number = int(value, 16) if value else None
            except ValueError:
                number = None
            if number is None or not 0 <= number < 0x80000000 or number in used:
                if value is not None:
                    pending.append(paragraph)
                continue
            used.add(number)
        candidate = 0x10000000
        for paragraph in pending:
            while candidate in used:
                candidate += 1
            paragraph.set(attribute, f"{candidate:08X}")
            used.add(candidate)
            candidate += 1
            result[local_name] += 1
    return result


def remove_empty_body_tables(document: DocumentClass) -> list[int]:
    """Remove empty revision shells left by accept/reject postprocessing."""
    removed = []
    for index, table in reversed(list(enumerate(document.tables))):
        if not _text(table._tbl).strip():
            parent = table._tbl.getparent()
            if parent is document.element.body:
                parent.remove(table._tbl)
                removed.append(index)
    return sorted(removed)


def apply_table_layout(document: DocumentClass, policy_value: dict[str, Any] | None = None) -> dict[str, Any]:
    policy = normalize_publication_layout_policy(policy_value)
    reports = []
    data_tables = 0
    diagram_tables = 0
    target_cells = 0
    checked_cells = 0
    target_paragraphs = 0
    checked_paragraphs = 0
    for table_index, table in enumerate(document.tables):
        records = _table_records(table)
        before_hash = _xml_sha256(table._tbl)
        diagram = _table_is_diagram(records, policy)
        if diagram:
            diagram_tables += 1
            reports.append({
                "table_index": table_index,
                "kind": "diagram",
                "cell_count": len(records),
                "nonzero_first_line_preserved": sum(
                    1 for record in records for p in record["tc"].findall(qn("w:p"))
                    if (p.find("./" + qn("w:pPr") + "/" + qn("w:ind")) is not None
                        and (p.find("./" + qn("w:pPr") + "/" + qn("w:ind")).get(qn("w:firstLine")) or "0") != "0")
                ),
                "sha256_before": before_hash,
                "sha256_after": before_hash,
                "preserved": True,
            })
            continue
        data_tables += 1
        semantics = _column_semantics(records, policy)
        font_pt = _table_font_pt(document, records)
        if not math.isfinite(font_pt) or font_pt <= 0:
            raise RuntimeError(f"table {table_index} has no effective font size")
        line_twips = _line_spacing_twips(font_pt, policy)
        table_target_paragraphs = 0
        for record in records:
            target_cells += 1
            classification = classify_cell_text(record["text"], policy)
            if classification == "diagram":
                raise RuntimeError(f"diagram cell escaped table-level classification in table {table_index}")
            alignment = (
                _empty_cell_alignment(record, semantics)
                if classification == "empty" else classification
            )
            paragraphs = record["tc"].findall(qn("w:p"))
            target_paragraphs += len(paragraphs)
            table_target_paragraphs += len(paragraphs)
            for paragraph in paragraphs:
                _set_paragraph_format(paragraph, alignment, line_twips)
                checked_paragraphs += 1
            _set_cell_margin(record["tc"], int(policy["table"]["cell_margin_twips"]))
            checked_cells += 1
        _set_row_flags(table, policy, diagram=False)
        reports.append({
            "table_index": table_index,
            "kind": "data",
            "cell_count": len(records),
            "paragraph_count": table_target_paragraphs,
            "effective_font_pt": font_pt,
            "line_spacing_pt": line_twips / 20.0,
            "sha256_before": before_hash,
            "sha256_after": _xml_sha256(table._tbl),
            "preserved": False,
        })
    if target_cells != checked_cells or target_paragraphs != checked_paragraphs:
        raise RuntimeError("publication table formatting coverage mismatch")
    return {
        "all_pass": True,
        "data_tables": data_tables,
        "diagram_tables": diagram_tables,
        "target_cells": target_cells,
        "checked_cells": checked_cells,
        "target_paragraphs": target_paragraphs,
        "checked_paragraphs": checked_paragraphs,
        "tables": reports,
    }


def _paragraph_format_state(paragraph) -> dict[str, Any]:
    ppr = paragraph.find(qn("w:pPr"))
    jc = ppr.find(qn("w:jc")) if ppr is not None else None
    ind = ppr.find(qn("w:ind")) if ppr is not None else None
    spacing = ppr.find(qn("w:spacing")) if ppr is not None else None
    return {
        "alignment": jc.get(qn("w:val")) if jc is not None else None,
        "indents": {
            name: (ind.get(qn(f"w:{name}")) if ind is not None else None)
            for name in ("left", "right", "start", "end", "firstLine", "hanging")
        },
        "line": spacing.get(qn("w:line")) if spacing is not None else None,
        "line_rule": spacing.get(qn("w:lineRule")) if spacing is not None else None,
        "before": spacing.get(qn("w:before")) if spacing is not None else None,
        "after": spacing.get(qn("w:after")) if spacing is not None else None,
    }


def audit_table_layout(document_or_path: DocumentClass | str | Path, policy_value: dict[str, Any] | None = None, *, math_line_rule: str = "exact") -> dict[str, Any]:
    if math_line_rule not in ("exact", "atLeast"):
        raise ValueError("Unsupported equation line rule")
    document = document_or_path if isinstance(document_or_path, DocumentClass) else Document(document_or_path)
    policy = normalize_publication_layout_policy(policy_value)
    target_tables: list[int] = []
    checked_tables: list[int] = []
    diagram_reports = []
    target_cells: set[tuple[int, int, int]] = set()
    checked_alignment: set[tuple[int, int, int]] = set()
    checked_indent: set[tuple[int, int, int]] = set()
    checked_spacing: set[tuple[int, int, int]] = set()
    failures: dict[str, list[dict[str, Any]]] = {
        "alignment": [], "indent": [], "spacing": [], "cell_margin": [],
        "cant_split": [], "repeat_header": [],
    }
    table_reports = []
    for table_index, table in enumerate(document.tables):
        records = _table_records(table)
        if _table_is_diagram(records, policy):
            diagram_reports.append({
                "table_index": table_index,
                "cell_count": len(records),
                "sha256": _xml_sha256(table._tbl),
                "nonzero_first_line_preserved": sum(
                    1 for record in records for p in record["tc"].findall(qn("w:p"))
                    if any(
                        value not in (None, "0")
                        for name, value in _paragraph_format_state(p)["indents"].items()
                        if name == "firstLine"
                    )
                ),
            })
            continue
        target_tables.append(table_index)
        semantics = _column_semantics(records, policy)
        font_pt = _table_font_pt(document, records)
        line_twips = _line_spacing_twips(font_pt, policy)
        for record in records:
            cell_key = (table_index, record["row"], record["cell"])
            target_cells.add(cell_key)
            classification = classify_cell_text(record["text"], policy)
            expected = _empty_cell_alignment(record, semantics) if classification == "empty" else classification
            cell_alignment_ok = True
            cell_indent_ok = True
            cell_spacing_ok = True
            for paragraph in record["tc"].findall(qn("w:p")):
                state = _paragraph_format_state(paragraph)
                # Word can remove a redundant direct w:jc on save. Resolve the
                # actual inherited value, not the presence of the XML element.
                from manuscript_format import effective_paragraph
                actual_alignment = effective_paragraph(document, paragraph).get("jc", {}).get("val", "left")
                if actual_alignment != expected:
                    cell_alignment_ok = False
                if any(value not in (None, "0") for value in state["indents"].values()):
                    cell_indent_ok = False
                if not (
                    state["line"] == str(line_twips)
                    and state["line_rule"] == (math_line_rule if paragraph.xpath('.//m:oMath') else "exact")
                    and state["before"] in (None, "0")
                    and state["after"] in (None, "0")
                ):
                    cell_spacing_ok = False
            if cell_alignment_ok:
                checked_alignment.add(cell_key)
            else:
                failures["alignment"].append({"table": table_index, "row": record["row"], "cell": record["cell"], "expected": expected})
            if cell_indent_ok:
                checked_indent.add(cell_key)
            else:
                failures["indent"].append({"table": table_index, "row": record["row"], "cell": record["cell"]})
            if cell_spacing_ok:
                checked_spacing.add(cell_key)
            else:
                failures["spacing"].append({"table": table_index, "row": record["row"], "cell": record["cell"], "expected_twips": line_twips})
            tcpr = record["tc"].find(qn("w:tcPr"))
            margins = tcpr.find(qn("w:tcMar")) if tcpr is not None else None
            expected_margin = str(int(policy["table"]["cell_margin_twips"]))
            margin_ok = margins is not None and all(
                (margins.find(qn(f"w:{side}")) is not None
                 and margins.find(qn(f"w:{side}")).get(qn("w:w")) == expected_margin)
                for side in ("left", "right")
            )
            if not margin_ok:
                failures["cell_margin"].append({"table": table_index, "row": record["row"], "cell": record["cell"]})
        rows = table._tbl.findall(qn("w:tr"))
        for row_index, tr in enumerate(rows):
            trpr = tr.find(qn("w:trPr"))
            if policy["table"]["cant_split_rows"] and (trpr is None or trpr.find(qn("w:cantSplit")) is None):
                failures["cant_split"].append({"table": table_index, "row": row_index})
        if len(rows) > 1 and policy["table"]["repeat_header"]:
            trpr = rows[0].find(qn("w:trPr"))
            if trpr is None or trpr.find(qn("w:tblHeader")) is None:
                failures["repeat_header"].append({"table": table_index})
        checked_tables.append(table_index)
        table_reports.append({
            "table_index": table_index,
            "cell_count": len(records),
            "effective_font_pt": font_pt,
            "line_spacing_pt": line_twips / 20.0,
        })
    coverage = {
        "target_tables": target_tables,
        "checked_tables": checked_tables,
        "target_cell_count": len(target_cells),
        "alignment_checked_count": len(checked_alignment),
        "indent_checked_count": len(checked_indent),
        "spacing_checked_count": len(checked_spacing),
    }
    coverage_ok = (
        bool(target_tables)
        and target_tables == checked_tables
        and target_cells == checked_alignment == checked_indent == checked_spacing
    )
    failed = [name for name, rows in failures.items() if rows]
    if not coverage_ok:
        failed.append("coverage")
    return {
        "all_pass": not failed,
        "coverage_ok": coverage_ok,
        "coverage": coverage,
        "failures": failures,
        "failed": failed,
        "data_tables": table_reports,
        "diagram_tables": diagram_reports,
        "historical_tables": [row for row in table_reports if row["table_index"] in {10, 11}],
    }


def usable_text_box_cm(document_or_path: DocumentClass | str | Path) -> list[dict[str, float]]:
    document = document_or_path if isinstance(document_or_path, DocumentClass) else Document(document_or_path)
    result = []
    for index, section in enumerate(document.sections):
        top = max(section.top_margin.cm, section.header_distance.cm)
        bottom = max(section.bottom_margin.cm, section.footer_distance.cm)
        result.append({
            "section": index,
            "width_cm": section.page_width.cm - section.left_margin.cm - section.right_margin.cm,
            "height_cm": section.page_height.cm - top - bottom,
            "top_constraint_cm": top,
            "bottom_constraint_cm": bottom,
        })
    return result


def audit_figure_geometry(document_or_path: DocumentClass | str | Path, policy_value: dict[str, Any] | None = None) -> dict[str, Any]:
    document = document_or_path if isinstance(document_or_path, DocumentClass) else Document(document_or_path)
    policy = normalize_publication_layout_policy(policy_value)
    boxes = usable_text_box_cm(document)
    if not boxes:
        raise RuntimeError("document has no section geometry")
    box = boxes[0]
    figure = policy["figure"]
    max_width = box["width_cm"] * float(figure["max_width_ratio"])
    max_height = (
        box["height_cm"]
        - float(figure["caption_block_cm"])
        - int(figure["narrative_lines"]) * float(figure["narrative_line_cm"])
    )
    shapes = []
    for index, shape in enumerate(document.inline_shapes):
        width = shape.width.cm
        height = shape.height.cm
        shapes.append({
            "index": index,
            "width_cm": width,
            "height_cm": height,
            "max_width_cm": max_width,
            "max_height_cm": max_height,
            "pass": width <= max_width + 0.01 and height <= max_height + 0.01,
        })
    return {
        "all_pass": bool(shapes) and all(row["pass"] for row in shapes),
        "sections": boxes,
        "shapes": shapes,
        "shape_count": len(shapes),
        "min_label_pt_advisory": float(figure["min_label_pt_advisory"]),
        "label_size_machine_verified": False,
    }


def _run(command: list[str], runner: Callable[[list[str]], Any] | None = None) -> str:
    if runner is not None:
        completed = runner(command)
        return completed.stdout
    return subprocess.run(command, check=True, text=True, capture_output=True).stdout


def _pdf_page_count(pdf: Path, runner=None) -> int:
    output = _run(["pdfinfo", str(pdf)], runner)
    for line in output.splitlines():
        if line.startswith("Pages:"):
            return int(line.split(":", 1)[1].strip())
    raise RuntimeError("pdfinfo did not report page count")


def _pdf_page_texts(pdf: Path, page_count: int, runner=None) -> dict[int, str]:
    result = {}
    for page in range(1, page_count + 1):
        result[page] = _run([
            "pdftotext", "-f", str(page), "-l", str(page), "-layout", str(pdf), "-"
        ], runner)
    return result


def _pdf_image_pages(pdf: Path, runner=None) -> list[int]:
    output = _run(["pdfimages", "-list", str(pdf)], runner)
    pages = set()
    for line in output.splitlines():
        match = re.match(r"\s*([0-9]+)\s+[0-9]+\s+", line)
        if match:
            pages.add(int(match.group(1)))
    return sorted(pages)


def _body_elements(document: DocumentClass):
    return list(document.element.body)


def _table_markers(document: DocumentClass, policy: dict[str, Any]) -> list[dict[str, Any]]:
    body = _body_elements(document)
    table_index = 0
    markers = []
    for body_index, element in enumerate(body):
        if element.tag != qn("w:tbl"):
            continue
        table = document.tables[table_index]
        records = _table_records(table)
        current_index = table_index
        table_index += 1
        if _table_is_diagram(records, policy):
            continue
        caption = ""
        for candidate in reversed(body[max(0, body_index - 3):body_index]):
            if candidate.tag == qn("w:p"):
                value = _text(candidate).strip()
                if CAPTION_RE.match(value):
                    caption = value
                    break
        header = [
            record["text"] for record in records
            if record["row"] == 0 and len(_normalized_text(record["text"])) >= 2
        ][:4]
        markers.append({"table_index": current_index, "caption": caption, "header": header})
    return markers


def _locate_table_pages(markers: list[dict[str, Any]], page_texts: dict[int, str]) -> tuple[dict[int, list[int]], list[int]]:
    normalized_pages = {page: _normalized_text(text) for page, text in page_texts.items()}
    located: dict[int, list[int]] = {}
    missing = []
    for marker in markers:
        caption_key = _normalized_text(marker["caption"])
        header_keys = [_normalized_text(value) for value in marker["header"] if _normalized_text(value)]
        pages = set()
        for page, text in normalized_pages.items():
            if caption_key and caption_key[:24] in text:
                pages.add(page)
            # Repeated headers locate continuation pages. Requiring every
            # available key avoids assigning unrelated tables to a page merely
            # because generic labels such as "变量" or "链条" recur.
            header_hits = sum(key in text for key in header_keys)
            if header_keys and header_hits == len(header_keys):
                pages.add(page)
        if pages:
            located[marker["table_index"]] = sorted(pages)
        else:
            missing.append(marker["table_index"])
    return located, missing


def _object_scope_pages(page_texts: dict[int, str], scope: dict[str, Any] | None) -> dict[str, Any]:
    if scope is None:
        pages = sorted(page_texts)
        return {"enabled": False, "all_pass": True, "pages": pages}
    start_key = _normalized_text(str(scope["start_text"]))
    end_key = _normalized_text(str(scope["end_text"]))
    normalized = {page: _normalized_text(text) for page, text in page_texts.items()}
    starts = [page for page, text in normalized.items() if start_key in text]
    if not starts:
        return {"enabled": True, "all_pass": False, "pages": [], "reason": "start_text_not_found"}
    start_page = min(starts)
    ends = [page for page, text in normalized.items() if page >= start_page and end_key in text]
    if not ends:
        return {"enabled": True, "all_pass": False, "pages": [], "reason": "end_text_not_found", "start_page": start_page}
    end_page = min(ends)
    if end_page < start_page:
        return {"enabled": True, "all_pass": False, "pages": [], "reason": "end_before_start"}
    include_end = bool(scope.get("include_end_page", True))
    final_page = end_page if include_end else max(start_page, end_page - 1)
    return {
        "enabled": True,
        "all_pass": True,
        "start_page": start_page,
        "end_page": end_page,
        "include_end_page": include_end,
        "pages": list(range(start_page, final_page + 1)),
    }


def _narrative_windows(document: DocumentClass, window: int = 20) -> list[str]:
    result = []
    seen = set()
    for element in _body_elements(document):
        if element.tag != qn("w:p") or element.xpath(".//w:drawing"):
            continue
        value = _text(element).strip()
        if not value or CAPTION_RE.match(value) or NOTE_RE.match(value) or HEADING_RE.match(value):
            continue
        normalized = _normalized_text(value)
        if len(normalized) < window * 2:
            continue
        for offset in range(0, len(normalized) - window + 1, window):
            token = normalized[offset:offset + window]
            if token not in seen:
                seen.add(token)
                result.append(token)
    return result


def evaluate_object_pages(
    object_pages: list[int] | set[int],
    narrative_characters: dict[int, int],
    *,
    minimum_characters: int = 40,
    minimum_lines: int = 2,
    window_size: int = 20,
) -> dict[str, Any]:
    rows = []
    for page in sorted(set(object_pages)):
        characters = int(narrative_characters.get(page, 0))
        lines = characters // max(1, window_size)
        passed = characters >= minimum_characters or lines >= minimum_lines
        rows.append({"page": page, "narrative_characters": characters, "narrative_lines_estimate": lines, "pass": passed})
    failed = [row["page"] for row in rows if not row["pass"]]
    return {"all_pass": not failed, "pages": rows, "object_only_pages": failed}


def audit_pdf_layout(
    document_path: str | Path,
    pdf_path: str | Path,
    policy_value: dict[str, Any] | None = None,
    *,
    runner: Callable[[list[str]], Any] | None = None,
) -> dict[str, Any]:
    document_path = Path(document_path)
    pdf_path = Path(pdf_path)
    document = Document(document_path)
    policy = normalize_publication_layout_policy(policy_value)
    page_count = _pdf_page_count(pdf_path, runner)
    page_texts = _pdf_page_texts(pdf_path, page_count, runner)
    image_pages = _pdf_image_pages(pdf_path, runner)
    markers = _table_markers(document, policy)
    table_pages, missing_tables = _locate_table_pages(markers, page_texts)
    all_table_pages = {page for pages in table_pages.values() for page in pages}
    all_object_pages = set(image_pages) | all_table_pages
    scope = _object_scope_pages(page_texts, policy["pagination"].get("object_scope"))
    scoped_pages = set(scope["pages"])
    object_pages = all_object_pages & scoped_pages
    windows = _narrative_windows(document)
    narrative_characters = {
        page: sum(len(token) for token in windows if token in _normalized_text(text))
        for page, text in page_texts.items()
    }
    pagination = policy["pagination"]
    page_gate = evaluate_object_pages(
        object_pages,
        narrative_characters,
        minimum_characters=int(pagination["minimum_narrative_characters"]),
        minimum_lines=int(pagination["minimum_narrative_lines"]),
    )
    figures = audit_figure_geometry(document, policy)
    inline_shape_count = len(document.inline_shapes)
    images_located = inline_shape_count == 0 or bool(image_pages)
    all_pass = (
        not missing_tables
        and images_located
        and scope["all_pass"]
        and page_gate["all_pass"]
        and figures["all_pass"]
    )
    return {
        "all_pass": all_pass,
        "page_count": page_count,
        "image_pages": image_pages,
        "inline_shape_count": inline_shape_count,
        "images_located": images_located,
        "table_pages": table_pages,
        "missing_table_indices": missing_tables,
        "object_pages_all": sorted(all_object_pages),
        "object_pages": sorted(object_pages),
        "object_scope": scope,
        "narrative_characters": narrative_characters,
        "object_page_gate": page_gate,
        "figure_geometry": figures,
    }


def write_json(path: str | Path, value: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
