import tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from native_table_fit import native_fit,transfer_geometry,signature
from journal_layout import normalize_publication_layout_policy

class NativeFitTests(unittest.TestCase):
 def test_content_default_and_explicit_override(self):
  self.assertEqual(normalize_publication_layout_policy({})["table"]["fit_mode"],"content")
  self.assertEqual(normalize_publication_layout_policy({"table":{"fit_mode":"fixed"}})["table"]["fit_mode"],"fixed")
  with self.assertRaises(ValueError):normalize_publication_layout_policy({"table":{"fit_mode":"fake"}})
 def test_command_is_native_and_preserve_skips(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/"owned.docx";d=Document();t=d.add_table(rows=2,cols=2)
   t.cell(0,0).text="变量";t.cell(0,1).text="95% CI"
   t.cell(1,0).text="AR-S";t.cell(1,1).text="-4.73 (-10.22, 0.77)"
   d.save(p);seen=[]
   def run(cmd,**kw):
    seen.append(cmd[2]);return SimpleNamespace(stdout="WORDREV_NATIVE_AUTOFIT_COMPLETE:content:1")
   q=native_fit(run,p,{})
   self.assertTrue(q["text_unchanged"])
   self.assertIn("auto fit behavior tableRef behavior auto fit content",seen[0])
   self.assertNotIn('activate',seen[0])
   n=len(seen);native_fit(run,p,{"table":{"fit_mode":"preserve"}});self.assertEqual(n,len(seen))
 def test_tracked_grid_and_cell_widths_keep_prior_snapshot(self):
  with tempfile.TemporaryDirectory() as td:
   a,b=Path(td)/"clean.docx",Path(td)/"red.docx"
   d=Document();t=d.add_table(rows=2,cols=2);t.cell(0,0).text="变量";t.cell(1,1).text="100"
   d.save(b);t.columns[0].width=Inches(1);t.cell(0,0).width=Inches(1);d.save(a)
   transfer_geometry(a,b,{},"Codex","2026-09-07T00:00:00Z")
   x=Document(b);self.assertTrue(x.element.xpath(".//w:tblGridChange"))
   for grid in x.element.xpath(".//w:tblGridChange"):
    self.assertEqual(set(grid.attrib),{qn('w:id')})
   self.assertTrue(x.element.xpath(".//w:tcPrChange"))
   self.assertEqual(signature(x.tables[0]),signature(Document(a).tables[0]))
   transfer_geometry(a,b,{},"Codex","2026-09-07T00:00:00Z")
   self.assertFalse(Document(b).element.xpath(".//w:tblGridChange//w:tblGridChange"))

if __name__=="__main__":unittest.main()
