import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

import thesis_format_adapter as a
from rendered_document_qa import check_page


class RenderingTests(unittest.TestCase):
    def test_all_story_indents_survive_save_and_reopen(self):
        d = Document()
        normal = d.styles["Normal"].element.get_or_add_pPr().get_or_add_ind()
        normal.set(qn("w:firstLine"), "420")
        normal.set(qn("w:firstLineChars"), "200")
        d.add_section(WD_SECTION_START.NEW_PAGE)
        d.sections[1].page_width, d.sections[1].page_height = d.sections[1].page_height, d.sections[1].page_width
        d.settings.odd_and_even_pages_header_footer = True
        for section in d.sections:
            section.different_first_page_header_footer = True
            for story in (section.header, section.even_page_header, section.first_page_header):
                story.is_linked_to_previous = False
                a._replace_story_text(story, "论文标题", font_ea="SimSun", font_latin="Times New Roman", size=10.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            for story in (section.footer, section.even_page_footer, section.first_page_footer):
                story.is_linked_to_previous = False
                a._replace_footer_with_page(story, font_latin="Times New Roman", size=10.5, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "fixture.docx"
            d.save(p)
            reopened = Document(p)
            for section in reopened.sections:
                for story in (section.header, section.even_page_header, section.first_page_header, section.footer, section.even_page_footer, section.first_page_footer):
                    ind = story.paragraphs[0]._p.pPr.ind
                    self.assertEqual(len(ind.attrib), 12)
                    self.assertTrue(all(v == "0" for v in ind.attrib.values()))

    def test_equation_namespace_and_pagination_ignore_but_token_structure_fail(self):
        ns = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
        before = etree.fromstring(f'<m:oMath xmlns:m="{ns}" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><m:r><m:t>x+1</m:t></m:r></m:oMath>')
        after = etree.fromstring(f'<n:oMath xmlns:n="{ns}" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><n:r><w:lastRenderedPageBreak/><n:t>x+1</n:t></n:r></n:oMath>')
        self.assertEqual(a._equation_projection(before), a._equation_projection(after))
        after[-1][-1].text = "x-1"
        self.assertNotEqual(a._equation_projection(before), a._equation_projection(after))
        after[-1][-1].text = "x+1"
        after[-1].tag = f'{{{ns}}}sSup'
        self.assertNotEqual(a._equation_projection(before), a._equation_projection(after))

    def test_formula_formatting_is_diagnostic_but_math_style_remains_semantic(self):
        ns = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
        wn = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        x = etree.fromstring(f'<m:oMath xmlns:m="{ns}" xmlns:w="{wn}"><m:r><m:t>x+1</m:t></m:r></m:oMath>')
        y = etree.fromstring(f'<m:oMath xmlns:m="{ns}" xmlns:w="{wn}"><m:r><w:rPr><w:sz w:val="24"/></w:rPr><m:t>x+1</m:t></m:r></m:oMath>')
        self.assertEqual(a._equation_projection(x), a._equation_projection(y))
        y[0].insert(0, etree.fromstring(f'<m:rPr xmlns:m="{ns}"><m:sty m:val="b"/></m:rPr>'))
        self.assertNotEqual(a._equation_projection(x), a._equation_projection(y))

    def test_native_paths_share_lock_and_pending_interlocks(self):
        import word_runtime
        with tempfile.TemporaryDirectory() as tmp, patch.object(Path, 'home', return_value=Path(tmp)):
            path = a._word_lock_path()
            self.assertEqual(path, word_runtime.word_pending_path().parent / 'word-automation.lock')
            with word_runtime.word_lock('fixture', timeout=1):
                with self.assertRaises(TimeoutError):
                    with a._word_process_lock(path, 0.05):
                        self.fail('two independent Word owners entered')
            pending = word_runtime.word_pending_path()
            pending.write_text('{"phase":"uncertain"}')
            with self.assertRaisesRegex(RuntimeError, 'WORD_RECOVERY_REQUIRED'):
                with a._word_process_lock(path, 1):
                    self.fail('pending state ignored')
    def test_theme_override_and_theme_regenerated_run_resolve_same_font(self):
        d = Document()
        run = d.add_paragraph().add_run("书目")
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsiaTheme"), "minorEastAsia")
        a._set_rfonts(run, "SimSun", "Times New Roman")
        self.assertNotIn(qn("w:eastAsiaTheme"), fonts.attrib)
        a.stabilize_theme_fonts(d, east_asia="SimSun", latin="Times New Roman")
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "font.docx"
            d.save(p)
            with ZipFile(p) as z:
                theme = etree.fromstring(z.read("word/theme/theme1.xml"))
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                self.assertEqual(theme.xpath('//a:minorFont/a:ea/@typeface', namespaces=ns), ["SimSun"])
                self.assertEqual(theme.xpath('//a:minorFont/a:font[@script="Hans"]/@typeface', namespaces=ns), ["SimSun"])

    def test_actual_glyph_font_contract_and_text_area_geometry(self):
        chars = [{"text": "中", "fontname": "ABCDEF+SimSun", "x0": 195, "x1": 205, "top": 20}]
        rule = {"font_regions": [{"box": [0, 0, 500, 40], "han": ["SimSun"]}],
                "stories": [{"text": "中", "box": [0, 0, 500, 40], "alignment": "center", "text_area": [50, 350], "tolerance_pt": 1}]}
        self.assertTrue(check_page(chars, rule)["pass"])
        chars[0]["fontname"] = "DengXian"
        self.assertFalse(check_page(chars, rule)["pass"])
        chars[0]["fontname"] = "SimSun"
        chars[0]["x0"] += 10.5
        chars[0]["x1"] += 10.5
        self.assertFalse(check_page(chars, rule)["pass"])

    def test_theme_language_changes_only_under_explicit_profile(self):
        d = Document()
        lang = d.settings.element.find(qn("w:themeFontLang"))
        if lang is None:
            lang = OxmlElement("w:themeFontLang")
            d.settings.element.append(lang)
        lang.set(qn("w:eastAsia"), "ja-JP")
        a.stabilize_theme_fonts(d, east_asia="SimSun", latin="Times New Roman")
        self.assertEqual(lang.get(qn("w:eastAsia")), "ja-JP")
        a.stabilize_theme_fonts(d, east_asia="SimSun", latin="Times New Roman", east_asia_language="zh-CN")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "font-language.docx"
            d.save(path)
            reopened = Document(path)
            self.assertEqual(reopened.settings.element.find(qn("w:themeFontLang")).get(qn("w:eastAsia")), "zh-CN")


if __name__ == "__main__":
    unittest.main()
