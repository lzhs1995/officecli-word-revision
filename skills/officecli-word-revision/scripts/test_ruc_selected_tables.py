"""Actual selected-table DOCX output, without invoking Word or OfficeCLI."""
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

import thesis_format_adapter as adapter

INDENTS = ("left", "right", "start", "end", "firstLine", "hanging",
           "leftChars", "rightChars", "startChars", "endChars", "firstLineChars", "hangingChars")
VALUES = ("收入", "12.30", "SC", "收入2020", "p < .05", "2020-09-19", "[0.1, 0.8]", "𠮷")


def set_bad_indents(element):
    ind = element.get_or_add_pPr().get_or_add_ind()
    for name in INDENTS:
        ind.set(qn("w:" + name), "200" if name.endswith("Chars") else "420")


class SelectedTableTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "fixture.docx"
        self.profile = json.loads((adapter.PROFILE_DIR / "ruc-doctoral-2026.json").read_text())

    def fixture(self, tables=1):
        doc = Document()
        set_bad_indents(doc.styles["Normal"].element)
        for _ in range(tables):
            table = doc.add_table(rows=1, cols=len(VALUES))
            for cell, value in zip(table.rows[0].cells, VALUES):
                cell.text = value
                p = cell.paragraphs[0]
                set_bad_indents(p._p)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return doc

    def apply(self, doc, policy="all_body_tables", **job):
        doc.save(self.path)
        adapter._apply_structural_formatting(self.path, {"scope": "chapter", "table_policy": policy, **job}, self.profile)
        return Document(self.path)

    def assert_zero(self, paragraph):
        ind = paragraph._p.pPr.find(qn("w:ind"))
        self.assertIsNotNone(ind)
        for name in INDENTS:
            self.assertEqual(ind.get(qn("w:" + name)), "0", name)

    def test_actual_saved_content_alignment_and_twelve_zero_indents(self):
        out = self.apply(self.fixture())
        expected = ("left", "right", "right", "left", "right", "right", "right", "left")
        for cell, value, alignment in zip(out.tables[0].rows[0].cells, VALUES, expected):
            self.assertEqual(cell.text, value)
            p = cell.paragraphs[0]
            self.assert_zero(p)
            self.assertEqual(p._p.pPr.find(qn("w:jc")).get(qn("w:val")), alignment)

    def test_inherited_indents_are_explicitly_overridden(self):
        doc = self.fixture()
        for cell in doc.tables[0].rows[0].cells:
            p = cell.paragraphs[0]._p
            p.pPr.remove(p.pPr.find(qn("w:ind")))
        for cell in self.apply(doc).tables[0].rows[0].cells:
            self.assert_zero(cell.paragraphs[0])

    def test_report_only_preserves_all_table_xml(self):
        doc = self.fixture(tables=2)
        before = [t._tbl.xml for t in doc.tables]
        self.assertEqual([t._tbl.xml for t in self.apply(doc, "report_only").tables], before)

    def test_explicit_only_changes_selected_top_level_table(self):
        doc = self.fixture(tables=2)
        before = doc.tables[0]._tbl.xml
        out = self.apply(doc, "explicit", table_indices=[2])
        self.assertEqual(out.tables[0]._tbl.xml, before)
        self.assert_zero(out.tables[1].cell(0, 0).paragraphs[0])

    def test_nested_and_merged_cells_preserve_structure_and_text(self):
        doc = self.fixture()
        table = doc.tables[0]
        table.cell(0, 0).merge(table.cell(0, 1))
        nested = table.cell(0, 2).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "3.140"
        p = nested.cell(0, 0).paragraphs[0]
        set_bad_indents(p._p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        before_text = table._tbl.xpath(".//w:t/text()")
        before_spans = table._tbl.xpath(".//w:gridSpan/@w:val")
        out = self.apply(doc).tables[0]
        self.assertEqual(out._tbl.xpath(".//w:t/text()"), before_text)
        self.assertEqual(out._tbl.xpath(".//w:gridSpan/@w:val"), before_spans)
        self.assertEqual(len(out._tbl.xpath(".//w:tbl")), 1)
        p = out.cell(0, 2).tables[0].cell(0, 0).paragraphs[0]
        self.assert_zero(p)
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_omml_content_is_preserved_and_right_aligned(self):
        doc = self.fixture()
        p = doc.tables[0].cell(0, 4).paragraphs[0]
        p.clear()
        math = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        token = OxmlElement("m:t")
        token.text = "β = .25"
        run.append(token); math.append(run); p._p.append(math)
        before = etree.tostring(math, method="c14n")
        out = self.apply(doc).tables[0].cell(0, 4).paragraphs[0]
        self.assertEqual(etree.tostring(out._p.find(qn("m:oMath")), method="c14n"), before)
        self.assertEqual(out.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assert_zero(out)

    def test_nested_cell_content_does_not_reclassify_parent_paragraphs(self):
        doc = self.fixture()
        cell = doc.tables[0].cell(0, 1)
        cell.text = "Parent"
        nested = cell.add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "子表"
        nested.cell(0, 1).text = "0.95"
        out = self.apply(doc).tables[0].cell(0, 1)
        for p in out.paragraphs:
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            self.assert_zero(p)
        self.assertEqual(out.tables[0].cell(0, 0).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(out.tables[0].cell(0, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_cell_with_han_in_later_paragraph_aligns_whole_cell_left(self):
        doc = self.fixture()
        doc.tables[0].cell(0, 1).add_paragraph("有效样本")
        for p in self.apply(doc).tables[0].cell(0, 1).paragraphs:
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assert_zero(p)

    def test_empty_cells_keep_alignment_but_clear_indents(self):
        doc = self.fixture()
        doc.tables[0].cell(0, 1).paragraphs[0].clear()
        p = self.apply(doc).tables[0].cell(0, 1).paragraphs[0]
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assert_zero(p)

    def test_other_profiles_without_opt_in_keep_paragraph_rules(self):
        for key in ("firstLineIndent", "hangingIndent", "hanAlignment", "mixedAlignment", "latinDigitAlignment"):
            del self.profile["table"][key]
        p = self.apply(self.fixture()).tables[0].cell(0, 0).paragraphs[0]
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(p._p.pPr.ind.get(qn("w:firstLineChars")), "200")

    def test_profile_alignment_values_are_consumed(self):
        self.profile["table"].update(hanAlignment="right", mixedAlignment="center", latinDigitAlignment="left")
        out = self.apply(self.fixture()).tables[0]
        self.assertEqual(out.cell(0, 0).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(out.cell(0, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(out.cell(0, 3).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_unsupported_indent_does_not_save_partial_changes(self):
        self.fixture().save(self.path)
        before = self.path.read_bytes()
        self.profile["table"]["firstLineIndent"] = "24pt"
        with self.assertRaisesRegex(ValueError, "explicit zero"):
            adapter._apply_structural_formatting(self.path, {"table_policy": "all_body_tables"}, self.profile)
        self.assertEqual(self.path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
